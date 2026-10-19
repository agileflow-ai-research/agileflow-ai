from io import BytesIO

import fitz
import pytest
from docx import Document

from app.main import app
from app.modules.agents.testing_validation.extraction import service as extraction_service_module
from app.modules.agents.testing_validation.inputs.router import get_agent4_input_storage
from app.modules.agents.testing_validation.inputs.storage import TestGenerationInputStorage
from app.modules.standalone.router import get_agent_project_repository, get_agent_run_repository
from test_standalone_foundation import FakeAgentProjectRepository, FakeAgentRunRepository


def pdf_bytes(*pages):
    document = fitz.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def docx_bytes():
    document = Document()
    document.add_heading("Checkout", level=1)
    document.add_paragraph("Customers pay with a saved card.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture
def extraction_context(authentication_context, tmp_path):
    projects = FakeAgentProjectRepository()
    runs = FakeAgentRunRepository()
    app.dependency_overrides[get_agent_project_repository] = lambda: projects
    app.dependency_overrides[get_agent_run_repository] = lambda: runs
    app.dependency_overrides[get_agent4_input_storage] = lambda: TestGenerationInputStorage(tmp_path)
    try:
        yield authentication_context, runs, tmp_path
    finally:
        for dependency in (get_agent_project_repository, get_agent_run_repository, get_agent4_input_storage):
            app.dependency_overrides.pop(dependency, None)


def create_project(client, agent_type="TESTING_VALIDATION", name="Quality Gate"):
    response = client.post("/api/v1/agent-projects", json={"agent_type": agent_type, "name": name, "domain": "Fintech"})
    assert response.status_code == 201
    return response.json()["data"]["agent_project_id"]


def base(project_id):
    return f"/api/v1/agents/testing-validation/{project_id}/test-generation/sources"


def upload(client, project_id, filename, payload, source_type="DOCUMENT", artifact_type="ACCEPTANCE_CRITERIA"):
    response = client.post(
        f"{base(project_id)}/files",
        data={"artifact_type": artifact_type, "source_type": source_type},
        files={"file": (filename, payload, "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def add_text(client, project_id, content="Refund within 14 days"):
    response = client.post(f"{base(project_id)}/text", json={"artifact_type": "ACCEPTANCE_CRITERIA", "content": content})
    assert response.status_code == 201
    return response.json()["data"]


def extract(client, project_id, source_id, force=False):
    suffix = "?force=true" if force else ""
    return client.post(f"{base(project_id)}/{source_id}/extract{suffix}")


def get_extraction(client, project_id, source_id):
    return client.get(f"{base(project_id)}/{source_id}/extraction")


@pytest.mark.parametrize("filename,payload,document_type,content_key", [
    ("criteria.pdf", pdf_bytes("Booking needs a code"), "PDF", "combined_text"),
    ("criteria.docx", docx_bytes(), "DOCX", "combined_text"),
    ("criteria.txt", b"Refund within 14 days", "TXT", "text"),
    ("criteria.md", b"# Refunds\n\nWithin 14 days", "MD", "markdown"),
    ("criteria.json", b'{"refund": {"days": 14}}', "JSON", "structured_content"),
    ("criteria.yaml", b"refund:\n  days: 14\n", "YAML", "structured_content"),
    ("criteria.yml", b"refund:\n  days: 14\n", "YAML", "structured_content"),
])
def test_each_supported_document_format_extracts_successfully(extraction_context, filename, payload, document_type, content_key):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, filename, payload)

    response = extract(client, project_id, source["source_id"])
    assert response.status_code == 200, response.text
    result = response.json()["data"]
    assert result["document_type"] == document_type
    assert result["status"] == "SUCCESS"
    assert content_key in result["content"]
    assert result["source_id"] == source["source_id"]
    assert result["source_sha256"] == source["sha256"]
    assert result["artifact_type"] == "ACCEPTANCE_CRITERIA"


def test_success_is_persisted_and_retrievable_with_status_on_the_source(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "a.txt", b"Refund within 14 days")
    created = extract(client, project_id, source["source_id"]).json()["data"]

    fetched = get_extraction(client, project_id, source["source_id"])
    assert fetched.status_code == 200
    assert fetched.json()["data"] == created
    listed = client.get(base(project_id)).json()["data"][0]
    assert listed["extraction_status"] == "SUCCESS"
    assert listed["processing_status"] == "UPLOADED"


def test_partial_extraction_is_preserved_not_reported_as_success(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "mixed.pdf", pdf_bytes("Has text", None))
    result = extract(client, project_id, source["source_id"]).json()["data"]
    assert result["status"] == "PARTIAL"
    assert get_extraction(client, project_id, source["source_id"]).json()["data"]["status"] == "PARTIAL"


def test_failed_extraction_is_persisted_safely_without_internal_detail(extraction_context, tmp_path):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "broken.pdf", b"%PDF-1.7 broken")
    response = extract(client, project_id, source["source_id"])
    assert response.status_code == 200
    result = response.json()["data"]
    assert result["status"] == "FAILED"
    assert "Traceback" not in response.text and "Error" not in result["warnings"][0]["message"]
    assert get_extraction(client, project_id, source["source_id"]).json()["data"]["status"] == "FAILED"


def test_extraction_source_must_belong_to_the_project_and_foreign_users_are_rejected(extraction_context):
    authentication, _, _ = extraction_context
    owner = authentication.client("USR-001")
    foreign = authentication.client("USR-002")
    project_a = create_project(owner, name="Project A")
    project_b = create_project(owner, name="Project B")
    source = upload(owner, project_a, "a.txt", b"Refund within 14 days")

    assert extract(owner, project_b, source["source_id"]).status_code == 404
    assert get_extraction(owner, project_b, source["source_id"]).status_code == 404
    assert extract(foreign, project_a, source["source_id"]).status_code == 404
    assert get_extraction(foreign, project_a, source["source_id"]).status_code == 404


def test_wrong_project_and_wrong_agent_type_are_rejected(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "a.txt", b"Refund within 14 days")
    requirements_project = create_project(client, agent_type="REQUIREMENTS")

    assert extract(client, "AGP-999", source["source_id"]).status_code == 404
    assert extract(client, requirements_project, source["source_id"]).status_code == 422
    assert get_extraction(client, requirements_project, source["source_id"]).status_code == 422


def test_image_and_text_sources_cannot_be_extracted(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    text = add_text(client, project_id)
    image = upload(client, project_id, "wire.png", b"\x89PNG\r\n\x1a\nminimal", source_type="IMAGE", artifact_type="WIREFRAMES")

    text_response = extract(client, project_id, text["source_id"])
    image_response = extract(client, project_id, image["source_id"])
    assert text_response.status_code == 422
    assert image_response.status_code == 422
    assert "pending analysis" in image_response.json()["errors"][0]["message"].lower()
    assert image["processing_status"] == "PENDING_ANALYSIS"


def test_unknown_source_id_is_rejected_without_leaking_storage(extraction_context):
    authentication, _, tmp_path = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    response = extract(client, project_id, "A4UP-" + "0" * 32)
    assert response.status_code == 404
    assert str(tmp_path) not in response.text


def test_extraction_responses_never_expose_filesystem_paths(extraction_context):
    authentication, _, tmp_path = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "criteria.pdf", pdf_bytes("Booking"))
    extracted = extract(client, project_id, source["source_id"])
    fetched = get_extraction(client, project_id, source["source_id"])
    listed = client.get(base(project_id))
    for response in (extracted, fetched, listed):
        assert str(tmp_path) not in response.text
        assert "/" not in source["source_id"] and "\\" not in response.text


def test_existing_success_is_reused_without_force_and_force_reruns(extraction_context, monkeypatch):
    authentication, _, tmp_path = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "a.txt", b"Refund within 14 days")
    first = extract(client, project_id, source["source_id"]).json()["data"]

    calls = []
    original = extraction_service_module.extract_document_safely

    def counted(*args, **kwargs):
        calls.append(args)
        return original(*args, **kwargs)

    monkeypatch.setattr(extraction_service_module, "extract_document_safely", counted)
    reused = extract(client, project_id, source["source_id"]).json()["data"]
    assert reused == first and calls == []

    rerun = extract(client, project_id, source["source_id"], force=True)
    assert rerun.status_code == 200
    assert len(calls) == 1
    sidecars = list(tmp_path.rglob(f"{source['source_id']}.extracted*"))
    assert len(sidecars) == 1


def test_changed_source_hash_invalidates_the_cached_extraction(extraction_context):
    authentication, _, tmp_path = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "a.txt", b"Refund within 14 days")
    extract(client, project_id, source["source_id"])
    sidecar = next(tmp_path.rglob(f"{source['source_id']}.extracted.json"))
    sidecar.write_text(sidecar.read_text(encoding="utf-8").replace(source["sha256"], "0" * 64), encoding="utf-8")

    assert get_extraction(client, project_id, source["source_id"]).status_code == 404
    assert extract(client, project_id, source["source_id"]).json()["data"]["source_sha256"] == source["sha256"]


def test_extraction_never_creates_an_agent_run(extraction_context):
    authentication, runs, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "a.pdf", pdf_bytes("Booking"))
    extract(client, project_id, source["source_id"])
    extract(client, project_id, source["source_id"], force=True)
    get_extraction(client, project_id, source["source_id"])
    assert runs.documents == {}


def test_archived_project_blocks_extraction_but_keeps_reads(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "a.txt", b"Refund within 14 days")
    extract(client, project_id, source["source_id"])
    client.patch(f"/api/v1/agent-projects/{project_id}", json={"status": "ARCHIVED"})
    assert extract(client, project_id, source["source_id"], force=True).status_code == 409
    assert get_extraction(client, project_id, source["source_id"]).status_code == 200


def test_multiple_document_sources_keep_independent_results(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project_id = create_project(client)
    good = upload(client, project_id, "good.txt", b"Refund within 14 days")
    partial = upload(client, project_id, "mixed.pdf", pdf_bytes("Has text", None))
    extract(client, project_id, good["source_id"])
    extract(client, project_id, partial["source_id"])

    statuses = {item["source_id"]: item["extraction_status"] for item in client.get(base(project_id)).json()["data"]}
    assert statuses == {good["source_id"]: "SUCCESS", partial["source_id"]: "PARTIAL"}
