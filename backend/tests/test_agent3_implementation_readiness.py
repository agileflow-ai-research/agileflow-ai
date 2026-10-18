import json
from io import BytesIO

import fitz
import pytest
from docx import Document

from app.main import app
from app.modules.agents.implementation.preparation.readiness.normalizer import (
    normalize_text_input,
)
from app.modules.agents.implementation.preparation.readiness.schemas import (
    FindingSeverity,
    ReadinessStatus,
    ReadinessTextInput,
)
from app.modules.agents.implementation.api.router import (
    get_github_repository_binding_repository,
    get_implementation_upload_storage,
)
from app.modules.agents.implementation.infrastructure.storage import ImplementationUploadStorage
from app.modules.standalone.router import (
    get_agent_project_repository,
    get_agent_run_repository,
)
from test_standalone_foundation import (
    FakeAgentProjectRepository,
    FakeAgentRunRepository,
)


class FakeBindingRepository:
    def __init__(self):
        self.documents = {}

    def find_by_project(self, project_id, owner_id):
        document = self.documents.get(project_id)
        if document and document["owner_id"] == owner_id:
            return document.copy()
        return None


@pytest.fixture
def readiness_context(authentication_context, tmp_path):
    projects = FakeAgentProjectRepository()
    runs = FakeAgentRunRepository()
    storage = ImplementationUploadStorage(tmp_path)
    bindings = FakeBindingRepository()
    app.dependency_overrides[get_agent_project_repository] = lambda: projects
    app.dependency_overrides[get_agent_run_repository] = lambda: runs
    app.dependency_overrides[get_implementation_upload_storage] = lambda: storage
    app.dependency_overrides[get_github_repository_binding_repository] = lambda: bindings
    try:
        yield authentication_context, projects, bindings, tmp_path
    finally:
        app.dependency_overrides.pop(get_agent_project_repository, None)
        app.dependency_overrides.pop(get_agent_run_repository, None)
        app.dependency_overrides.pop(get_implementation_upload_storage, None)
        app.dependency_overrides.pop(get_github_repository_binding_repository, None)


def create_project(client, agent_type="IMPLEMENTATION"):
    response = client.post(
        "/api/v1/agent-projects",
        json={"agent_type": agent_type, "name": "Builder", "domain": "SaaS"},
    )
    assert response.status_code == 201
    return response.json()["data"]


def upload(client, project_id, filename, content, *, source_type="DOCUMENT", artifact="API_SPECIFICATION"):
    response = client.post(
        f"/api/v1/agents/implementation/{project_id}/input-files",
        data={"source_type": source_type, "artifact_type": artifact},
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert response.status_code == 201
    return response.json()["data"]


def extract(client, project_id, upload_id):
    response = client.post(
        f"/api/v1/agents/implementation/{project_id}/input-files/{upload_id}/extract"
    )
    assert response.status_code == 200
    return response.json()["data"]


def evaluate(client, project_id, payload):
    return client.post(
        f"/api/v1/agents/implementation/{project_id}/readiness/evaluate",
        json=payload,
    )


def text_input(artifact="FUNCTIONAL_REQUIREMENTS", content="Users can create bookings.", input_id="LOCAL-1"):
    return {
        "input_id": input_id,
        "artifact_type": artifact,
        "content": content,
    }


def payload(*, outputs=("BACKEND_CODE",), texts=None, documents=None, images=None, **extra):
    value = {
        "selected_outputs": list(outputs),
        "text_inputs": texts if texts is not None else [text_input()],
        "document_upload_ids": documents or [],
        "image_upload_ids": images or [],
        "human_resolutions": [],
    }
    value.update(extra)
    return value


def pdf_bytes(*pages):
    document = fitz.open()
    for value in pages:
        page = document.new_page()
        if value:
            page.insert_text((72, 72), value)
    result = document.tobytes()
    document.close()
    return result


def docx_bytes():
    document = Document()
    document.add_heading("Architecture", 1)
    document.add_paragraph("Backend: Express.js")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def finding_codes(result):
    return {item["code"] for item in result["findings"]}


def test_exact_readiness_enums_and_text_normalization():
    assert {item.value for item in ReadinessStatus} == {
        "READY", "READY_WITH_WARNINGS", "NOT_READY"
    }
    assert {item.value for item in FindingSeverity} == {"ERROR", "WARNING", "INFO"}
    item = ReadinessTextInput(
        input_id="LOCAL-1", artifact_type="FUNCTIONAL_REQUIREMENTS",
        label=" Requirements ", content="  line one\nline two  ",
    )
    normalized = normalize_text_input(item)
    assert normalized.text_content == "line one\nline two"
    assert normalized.label == "Requirements"


def test_no_output_no_input_and_no_behavior_are_not_ready(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    no_context = evaluate(client, project["agent_project_id"], payload(outputs=(), texts=[])).json()["data"]
    assert no_context["status"] == "NOT_READY"
    assert finding_codes(no_context) >= {
        "NO_OUTPUT_SELECTED", "NO_IMPLEMENTATION_CONTEXT", "IMPLEMENTATION_BEHAVIOR_MISSING"
    }
    tech_only = evaluate(
        client, project["agent_project_id"],
        payload(outputs=("PROJECT_STRUCTURE",), texts=[text_input("TECHNOLOGY_STACK", "Backend: Express")]),
    ).json()["data"]
    assert "IMPLEMENTATION_BEHAVIOR_MISSING" in finding_codes(tech_only)


def test_extraction_missing_failed_partial_and_success(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    project_id = project["agent_project_id"]
    missing = upload(client, project_id, "missing.txt", b"API spec")
    result = evaluate(client, project_id, payload(documents=[missing["upload_id"]])).json()["data"]
    assert "DOCUMENT_NOT_EXTRACTED" in finding_codes(result)

    failed = upload(client, project_id, "empty.pdf", pdf_bytes(""))
    extract(client, project_id, failed["upload_id"])
    result = evaluate(client, project_id, payload(documents=[failed["upload_id"]])).json()["data"]
    assert "DOCUMENT_EXTRACTION_FAILED" in finding_codes(result)

    partial = upload(client, project_id, "partial.pdf", pdf_bytes("API", ""))
    extract(client, project_id, partial["upload_id"])
    result = evaluate(client, project_id, payload(documents=[partial["upload_id"]])).json()["data"]
    assert "DOCUMENT_EXTRACTION_PARTIAL" in finding_codes(result)

    success = upload(client, project_id, "api.pdf", pdf_bytes("API"))
    extract(client, project_id, success["upload_id"])
    result = evaluate(client, project_id, payload(documents=[success["upload_id"]])).json()["data"]
    assert result["artifact_manifest"][-1]["availability"] == "AVAILABLE"


@pytest.mark.parametrize(
    ("output", "code"),
    [
        ("FRONTEND_CODE", "FRONTEND_UI_CONTEXT_LIMITED"),
        ("BACKEND_CODE", "BACKEND_API_CONTEXT_LIMITED"),
        ("DATABASE_IMPLEMENTATION", "DATABASE_SCHEMA_REQUIRED"),
        ("API_IMPLEMENTATION", "API_SPECIFICATION_REQUIRED"),
        ("CONTAINER_CONFIGURATION", "CONTAINER_CONTEXT_DEFAULTS_REQUIRED"),
        ("GIT_SUMMARY", "GITHUB_REPOSITORY_REQUIRED"),
        ("PROJECT_STRUCTURE", "PROJECT_STRUCTURE_TECHNICAL_CONTEXT_LIMITED"),
    ],
)
def test_output_aware_rules(readiness_context, output, code):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    result = evaluate(client, project["agent_project_id"], payload(outputs=(output,))).json()["data"]
    assert code in finding_codes(result)


def test_context_warnings_info_and_ready_statuses(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    project_id = project["agent_project_id"]
    warning = evaluate(client, project_id, payload(outputs=("CONFIGURATION_FILES",))).json()["data"]
    assert warning["status"] == "READY_WITH_WARNINGS"
    assert finding_codes(warning) >= {"PRECODE_TEST_CONTEXT_MISSING", "PROTOTYPE_STACK_DEFAULT_APPLIED"}

    ready = evaluate(
        client,
        project_id,
        payload(
            outputs=("PROJECT_STRUCTURE",),
            texts=[
                text_input(),
                text_input("TECHNOLOGY_STACK", "Frontend: React\nBackend: Express.js\nDatabase: MySQL", "LOCAL-2"),
                text_input("UNIT_TEST_CASES", "Test booking creation", "LOCAL-3"),
            ],
        ),
    ).json()["data"]
    assert ready["status"] == "READY"
    assert ready["effective_stack"] == {
        "frontend": "React + Vite", "styling": "Plain CSS",
        "backend": "Node.js + Express.js", "database": "MySQL",
        "language": "JavaScript", "api_style": "REST",
    }


def test_image_is_pending_not_understood(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    image = upload(
        client, project["agent_project_id"], "wireframe.png",
        b"\x89PNG\r\n\x1a\nminimal", source_type="IMAGE", artifact="UI_WIREFRAMES",
    )
    result = evaluate(
        client, project["agent_project_id"],
        payload(outputs=("FRONTEND_CODE",), images=[image["upload_id"]]),
    ).json()["data"]
    assert finding_codes(result) >= {"VISUAL_CONTEXT_PENDING", "UI_WIREFRAME_ANALYSIS_PENDING"}
    assert result["coverage"]["visual_context_pending"] is True
    assert result["artifact_manifest"][-1]["availability"] == "PENDING_ANALYSIS"


@pytest.mark.parametrize(
    ("declaration", "category"),
    [
        ("Frontend: Angular", "frontend"),
        ("Backend: FastAPI", "backend"),
        ("Database: PostgreSQL", "database"),
        ("Database: MongoDB", "database"),
    ],
)
def test_explicit_unsupported_stack_detection_and_resolution(readiness_context, declaration, category):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    request = payload(
        outputs=("CONFIGURATION_FILES",),
        texts=[text_input(content=f"Users can book.\n{declaration}")],
    )
    first = evaluate(client, project["agent_project_id"], request).json()["data"]
    conflict = next(item for item in first["findings"] if item["code"] == "UNSUPPORTED_PROTOTYPE_TECHNOLOGY")
    assert first["status"] == "NOT_READY"
    assert conflict["category"] == category
    assert conflict["requires_human_decision"] is True
    request.update({
        "input_fingerprint": first["input_fingerprint"],
        "human_resolutions": [{
            "finding_id": conflict["finding_id"],
            "decision": "ADAPT_TO_SUPPORTED_STACK",
            "note": "Approved for PP1 prototype.",
        }],
    })
    resolved = evaluate(client, project["agent_project_id"], request).json()["data"]
    finding = next(item for item in resolved["findings"] if item["finding_id"] == conflict["finding_id"])
    assert finding["resolved"] is True
    assert resolved["status"] == "READY_WITH_WARNINGS"


def test_incidental_or_negated_technology_is_not_a_conflict(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    for content in ("FastAPI was considered historically.", "Database: Do not use MongoDB"):
        result = evaluate(
            client, project["agent_project_id"],
            payload(outputs=("CONFIGURATION_FILES",), texts=[text_input(content=content)]),
        ).json()["data"]
        assert "UNSUPPORTED_PROTOTYPE_TECHNOLOGY" not in finding_codes(result)


def test_structured_json_yaml_and_docx_are_loaded_from_trusted_extraction(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    project_id = project["agent_project_id"]
    uploads = [
        upload(client, project_id, "stack.json", b'{"backend": "FastAPI"}', artifact="TECHNOLOGY_STACK"),
        upload(client, project_id, "stack.yaml", b"database: MongoDB\n", artifact="TECHNOLOGY_STACK"),
        upload(client, project_id, "architecture.docx", docx_bytes(), artifact="SYSTEM_ARCHITECTURE"),
    ]
    for item in uploads:
        extract(client, project_id, item["upload_id"])
    result = evaluate(
        client, project_id,
        payload(outputs=("CONFIGURATION_FILES",), documents=[item["upload_id"] for item in uploads]),
    ).json()["data"]
    conflicts = [item for item in result["findings"] if item["code"] == "UNSUPPORTED_PROTOTYPE_TECHNOLOGY"]
    assert {item["category"] for item in conflicts} == {"backend", "database"}
    assert "content" not in json.dumps(result["artifact_manifest"]).lower()
    assert '"path"' not in json.dumps(result).lower()


def test_cross_source_conflict_is_deterministic(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    request = payload(
        outputs=("CONFIGURATION_FILES",),
        texts=[
            text_input(content="Users can book.\nDatabase: MySQL", input_id="LOCAL-A"),
            text_input("SYSTEM_ARCHITECTURE", "Database: MongoDB", "LOCAL-B"),
        ],
    )
    first = evaluate(client, project["agent_project_id"], request).json()["data"]
    second = evaluate(client, project["agent_project_id"], request).json()["data"]
    assert first["input_fingerprint"] == second["input_fingerprint"]
    assert "CROSS_SOURCE_TECHNOLOGY_CONFLICT" in finding_codes(first)


def test_invalid_unknown_and_stale_resolutions_rejected(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    request = payload(outputs=("CONFIGURATION_FILES",))
    first = evaluate(client, project["agent_project_id"], request).json()["data"]
    request["input_fingerprint"] = first["input_fingerprint"]
    request["human_resolutions"] = [{
        "finding_id": "RF-0000000000000000", "decision": "ADAPT_TO_SUPPORTED_STACK"
    }]
    assert evaluate(client, project["agent_project_id"], request).status_code == 422

    request["text_inputs"][0]["content"] = "Changed behavior"
    assert evaluate(client, project["agent_project_id"], request).status_code == 409


def test_fingerprint_changes_with_text_output_and_document_sha(readiness_context):
    authentication, _, _, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    project_id = project["agent_project_id"]
    base = payload(outputs=("CONFIGURATION_FILES",))
    first = evaluate(client, project_id, base).json()["data"]["input_fingerprint"]
    changed_text = payload(outputs=("CONFIGURATION_FILES",), texts=[text_input(content="Changed")])
    second = evaluate(client, project_id, changed_text).json()["data"]["input_fingerprint"]
    changed_output = payload(outputs=("PROJECT_STRUCTURE",))
    third = evaluate(client, project_id, changed_output).json()["data"]["input_fingerprint"]
    document = upload(client, project_id, "api.txt", b"API", artifact="API_SPECIFICATION")
    extract(client, project_id, document["upload_id"])
    with_document = payload(outputs=("CONFIGURATION_FILES",), documents=[document["upload_id"]])
    fourth = evaluate(client, project_id, with_document).json()["data"]["input_fingerprint"]
    assert len({first, second, third, fourth}) == 4


def test_repository_binding_satisfies_git_summary(readiness_context):
    authentication, _, bindings, _ = readiness_context
    client = authentication.client()
    project = create_project(client)
    bindings.documents[project["agent_project_id"]] = {
        "owner_id": "USR-001", "repository_id": 123
    }
    result = evaluate(
        client, project["agent_project_id"], payload(outputs=("GIT_SUMMARY",))
    ).json()["data"]
    assert "GITHUB_REPOSITORY_REQUIRED" not in finding_codes(result)


def test_project_upload_security_guards(readiness_context):
    authentication, _, _, _ = readiness_context
    owner = authentication.client("USR-001")
    foreign = authentication.client("USR-002")
    project = create_project(owner)
    document = upload(owner, project["agent_project_id"], "api.txt", b"API")
    assert evaluate(
        foreign, project["agent_project_id"], payload(documents=[document["upload_id"]])
    ).status_code == 404
    owner.patch(
        f'/api/v1/agent-projects/{project["agent_project_id"]}',
        json={"status": "ARCHIVED"},
    )
    assert evaluate(owner, project["agent_project_id"], payload()).status_code == 409
    wrong = create_project(owner, "DESIGN_PLANNING")
    assert evaluate(owner, wrong["agent_project_id"], payload()).status_code == 422
