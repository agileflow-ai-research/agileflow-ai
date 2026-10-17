import docx
import pypdf
import pytest

from app.common.transcription_provider import (
    OpenAICompatibleSpeechToTextProvider,
    TranscriptionProviderError,
    TranscriptionProviderConfigurationError,
    build_transcription_provider,
)
from app.core.config import Settings
from app.main import app
from app.modules.agent1.router import get_agent1_engine, get_speech_to_text_provider
from app.modules.agent1.router import get_requirement_source_processing_service
from app.modules.agent1.processing import (
    AudioProcessor,
    DocumentProcessor,
    FFmpegAudioExtractor,
    NoUsableAudioError,
    SourceNormalizer,
    SourceProcessorResolver,
    VideoAudioExtractionError,
    VideoProcessor,
    VideoProcessorNotConfiguredError,
)
from app.modules.agent1.schemas import SourceProcessingStatus
from app.modules.agent1.service import RequirementSourceProcessingService
from app.modules.standalone.router import get_standalone_service
from test_agent1_sources import create_project, sources_context


class FakeSpeechToTextProvider:
    def __init__(self, transcripts=None, error=None):
        self.transcripts = transcripts or {}
        self.error = error
        self.calls = []

    def transcribe(self, file_path):
        self.calls.append(file_path)
        if self.error:
            raise self.error
        value = self.transcripts.get(file_path.name)
        if value is None:
            value = next(
                (transcript for name, transcript in self.transcripts.items() if file_path.name.endswith(name)),
                self.transcripts.get("*", ""),
            )
        if isinstance(value, Exception):
            raise value
        return value


class FakeVideoAudioExtractor:
    def __init__(self, error=None):
        self.error = error
        self.output_paths = []

    def extract(self, video_path, output_path):
        self.output_paths.append(output_path)
        if self.error:
            raise self.error
        output_path.write_bytes(b"extracted audio")


def configure_video_processing(projects, runs, sources, provider, extractor):
    def override():
        audio_processor = AudioProcessor(provider)
        resolver = SourceProcessorResolver(
            audio_processor=audio_processor,
            video_processor=VideoProcessor(extractor, audio_processor),
        )
        return RequirementSourceProcessingService(
            get_standalone_service(projects, runs), sources, resolver
        )

    app.dependency_overrides[get_requirement_source_processing_service] = override


@pytest.fixture
def configured_transcription_provider():
    provider = FakeSpeechToTextProvider({"*": "  Customer interview\r\n\r\n\r\nBooking details  "})
    app.dependency_overrides[get_speech_to_text_provider] = lambda: provider
    try:
        yield provider
    finally:
        app.dependency_overrides.pop(get_speech_to_text_provider, None)


def test_txt_is_normalized(tmp_path):
    path = tmp_path / "requirements.txt"
    path.write_bytes(b"  First requirement  \r\n\r\n\r\nSecond requirement  ")
    assert DocumentProcessor().process(path) == "First requirement\n\nSecond requirement"


def test_pdf_machine_text_is_normalized(tmp_path, monkeypatch):
    path = tmp_path / "requirements.pdf"
    path.write_bytes(b"%PDF fake")

    class Page:
        def extract_text(self):
            return "PDF requirement  \r\n\r\n\r\nMore detail"

    class Reader:
        pages = [Page()]

    monkeypatch.setattr(pypdf, "PdfReader", lambda stream: Reader())
    assert DocumentProcessor().process(path) == "PDF requirement\n\nMore detail"


def test_docx_paragraphs_and_tables_are_extracted(tmp_path):
    document = docx.Document()
    document.add_paragraph("Booking requirement")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Actor"
    table.cell(0, 1).text = "Librarian"
    path = tmp_path / "requirements.docx"
    document.save(path)
    text = DocumentProcessor().process(path)
    assert text == "Booking requirement\nActor\tLibrarian"


def test_normalizer_rejects_effectively_empty_content():
    try:
        SourceNormalizer().normalize(" \r\n\r\n ")
    except ValueError as exc:
        assert str(exc) == "No usable text could be extracted."
    else:
        raise AssertionError("empty normalized content was accepted")


def test_openai_compatible_transcription_provider_uses_audio_endpoint(tmp_path, monkeypatch):
    audio = tmp_path / "interview.mp3"
    audio.write_bytes(b"audio")
    captured = {}

    class Response:
        def raise_for_status(self):
            return None

        def json(self):
            return {"text": "Interview transcript"}

    def fake_post(url, **kwargs):
        captured["url"] = url
        captured.update(kwargs)
        return Response()

    monkeypatch.setattr("app.common.transcription_provider.httpx.post", fake_post)
    provider = OpenAICompatibleSpeechToTextProvider(
        api_key="test-key", model="whisper-1", base_url="https://speech.example/v1",
        timeout_seconds=30, language="en",
    )
    assert provider.transcribe(audio) == "Interview transcript"
    assert captured["url"] == "https://speech.example/v1/audio/transcriptions"
    assert captured["data"] == {"model": "whisper-1", "response_format": "json", "language": "en"}
    assert captured["headers"] == {"Authorization": "Bearer test-key"}


def test_transcription_provider_configuration_is_independent_and_safe(tmp_path):
    provider = build_transcription_provider(Settings(
        mongodb_uri="mongodb://localhost:27017",
        model_api_key="generation-key",
        transcription_api_key=None,
    ))
    with pytest.raises(TranscriptionProviderConfigurationError) as exc_info:
        provider.transcribe(tmp_path / "audio.mp3")
    assert "TRANSCRIPTION_API_KEY" in str(exc_info.value)
    assert "generation-key" not in str(exc_info.value)


def test_processing_endpoint_completes_document_without_creating_run(sources_context):
    auth, _, runs, sources, _ = sources_context
    client = auth.client("USR-001")
    project = create_project(client)
    upload = client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
        files=[("files", ("requirements.txt", b"First\r\n\r\n\r\nSecond requirement", "text/plain"))],
    )
    source_id = upload.json()["data"][0]["source_id"]

    response = client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
    )
    assert response.status_code == 200
    assert response.json()["data"] == [{
        "source_id": source_id,
        "source_type": "DOCUMENT",
        "processing_status": "COMPLETED",
        "processing_error_code": None,
        "processing_error_message": None,
    }]
    persisted = sources.documents[(project["agent_project_id"], source_id)]
    assert persisted["normalized_text"] == "First\n\nSecond requirement"
    assert persisted["processing_status"] == "COMPLETED"
    assert persisted["processed_at"] is not None
    assert sources.processing_updates == [
        (source_id, SourceProcessingStatus.PROCESSING),
        (source_id, SourceProcessingStatus.COMPLETED),
    ]
    assert runs.documents == {}
    assert "storage_path" not in response.text


def test_one_document_failure_does_not_fail_other_documents(sources_context):
    auth, _, _, sources, _ = sources_context
    client = auth.client("USR-001")
    project = create_project(client)
    client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
        files=[
            ("files", ("good.txt", b"A complete useful requirement", "text/plain")),
            ("files", ("empty.txt", b"   ", "text/plain")),
        ],
    )
    response = client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
    )
    data = response.json()["data"]
    assert [item["processing_status"] for item in data] == ["COMPLETED", "FAILED"]
    assert data[1]["processing_error_code"] == "NO_EXTRACTABLE_TEXT"
    assert data[1]["processing_error_message"] == "No usable text could be extracted."
    assert "D:\\" not in response.text and "Traceback" not in response.text
    assert sources.documents[(project["agent_project_id"], "DOC-001")]["normalized_text"]


@pytest.mark.parametrize(
    ("filename", "mime_type"),
    [("walkthrough.mp4", "video/mp4"), ("walkthrough.mov", "video/quicktime"), ("walkthrough.webm", "video/webm")],
)
def test_video_formats_extract_delegate_normalize_and_cleanup(
    sources_context, filename, mime_type
):
    auth, projects, runs, sources, _ = sources_context
    provider = FakeSpeechToTextProvider({"*": "  Video interview\r\n\r\n\r\nBooking flow  "})
    extractor = FakeVideoAudioExtractor()
    configure_video_processing(projects, runs, sources, provider, extractor)
    try:
        client = auth.client("USR-001")
        project = create_project(client)
        upload = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[("files", (filename, b"video bytes", mime_type))],
        ).json()["data"][0]
        assert upload["source_id"] == "VID-001"
        result = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        ).json()["data"][0]
        assert result["source_id"] == "VID-001"
        assert result["processing_status"] == "COMPLETED"
        persisted = sources.documents[(project["agent_project_id"], "VID-001")]
        assert persisted["normalized_text"] == "Video interview\n\nBooking flow"
        assert sources.processing_updates == [
            ("VID-001", SourceProcessingStatus.PROCESSING),
            ("VID-001", SourceProcessingStatus.COMPLETED),
        ]
        assert provider.calls[0].suffix == ".wav"
        assert extractor.output_paths and not extractor.output_paths[0].parent.exists()
        assert runs.documents == {}
    finally:
        app.dependency_overrides.pop(get_requirement_source_processing_service, None)


@pytest.mark.parametrize(
    ("filename", "mime_type"),
    [("interview.mp3", "audio/mpeg"), ("interview.wav", "audio/wav"), ("interview.m4a", "audio/mp4")],
)
def test_audio_formats_transcribe_normalize_and_preserve_traceability(
    sources_context, configured_transcription_provider, filename, mime_type
):
    auth, _, runs, sources, _ = sources_context
    client = auth.client("USR-001")
    project = create_project(client)
    upload = client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
        files=[("files", (filename, b"audio bytes", mime_type))],
    ).json()["data"][0]
    assert upload["source_id"] == "AUD-001"

    result = client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
    ).json()["data"][0]
    assert result["source_id"] == "AUD-001"
    assert result["processing_status"] == "COMPLETED"
    persisted = sources.documents[(project["agent_project_id"], "AUD-001")]
    assert persisted["normalized_text"] == "Customer interview\n\nBooking details"
    assert sources.processing_updates == [
        ("AUD-001", SourceProcessingStatus.PROCESSING),
        ("AUD-001", SourceProcessingStatus.COMPLETED),
    ]
    assert runs.documents == {}


def test_blank_audio_transcript_fails_safely(sources_context):
    provider = FakeSpeechToTextProvider({"*": "  \r\n "})
    app.dependency_overrides[get_speech_to_text_provider] = lambda: provider
    try:
        auth, _, _, sources, _ = sources_context
        client = auth.client("USR-001")
        project = create_project(client)
        client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[("files", ("blank.wav", b"audio bytes", "audio/wav"))],
        )
        result = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        ).json()["data"][0]
        assert result["processing_status"] == "FAILED"
        assert result["processing_error_code"] == "NO_USABLE_TRANSCRIPT"
        assert result["processing_error_message"] == "No usable transcript could be produced."
        assert sources.documents[(project["agent_project_id"], "AUD-001")]["processed_at"]
    finally:
        app.dependency_overrides.pop(get_speech_to_text_provider, None)


def test_provider_failure_is_sanitized_and_other_sources_continue(sources_context):
    provider = FakeSpeechToTextProvider({
        "bad.mp3": TranscriptionProviderError("secret-token C:\\provider\\internal"),
        "good.wav": "A valid stakeholder interview transcript",
    })
    app.dependency_overrides[get_speech_to_text_provider] = lambda: provider
    try:
        auth, _, runs, sources, _ = sources_context
        client = auth.client("USR-001")
        project = create_project(client)
        client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[
                ("files", ("bad.mp3", b"bad audio", "audio/mpeg")),
                ("files", ("good.wav", b"good audio", "audio/wav")),
            ],
        )
        response = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        )
        data = response.json()["data"]
        assert [item["processing_status"] for item in data] == ["FAILED", "COMPLETED"]
        assert data[0]["processing_error_code"] == "AUDIO_TRANSCRIPTION_FAILED"
        assert data[0]["processing_error_message"] == "The audio source could not be transcribed."
        assert "secret-token" not in response.text and "provider\\internal" not in response.text
        assert sources.documents[(project["agent_project_id"], "AUD-002")]["normalized_text"]
        assert runs.documents == {}
    finally:
        app.dependency_overrides.pop(get_speech_to_text_provider, None)


def test_missing_transcription_config_fails_audio_without_invoking_agent_engine(sources_context):
    app.dependency_overrides[get_agent1_engine] = lambda: (_ for _ in ()).throw(
        AssertionError("Agent1Engine must not be resolved")
    )
    try:
        auth, _, runs, _, _ = sources_context
        client = auth.client("USR-001")
        project = create_project(client)
        client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[("files", ("interview.mp3", b"audio", "audio/mpeg"))],
        )
        result = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        ).json()["data"][0]
        assert result["processing_status"] == "FAILED"
        assert result["processing_error_code"] == "TRANSCRIPTION_PROVIDER_NOT_CONFIGURED"
        assert result["processing_error_message"] == "Audio transcription is not configured."
        assert runs.documents == {}
    finally:
        app.dependency_overrides.pop(get_agent1_engine, None)


def test_mixed_document_and_audio_processing(sources_context, configured_transcription_provider):
    auth, _, _, sources, _ = sources_context
    client = auth.client("USR-001")
    project = create_project(client)
    client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
        files=[
            ("files", ("brief.txt", b"A useful document requirement", "text/plain")),
            ("files", ("interview.m4a", b"audio bytes", "audio/mp4")),
        ],
    )
    data = client.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
    ).json()["data"]
    assert [(item["source_id"], item["processing_status"]) for item in data] == [
        ("DOC-001", "COMPLETED"), ("AUD-001", "COMPLETED")
    ]
    assert sources.documents[(project["agent_project_id"], "DOC-001")]["normalized_text"]
    assert sources.documents[(project["agent_project_id"], "AUD-001")]["normalized_text"]


def test_ffmpeg_uses_shell_free_arguments_and_reports_missing_binary(tmp_path, monkeypatch):
    video = tmp_path / "video.mp4"
    video.write_bytes(b"video")
    output = tmp_path / "audio.wav"
    captured = {}

    def missing(arguments, **kwargs):
        captured["arguments"] = arguments
        captured.update(kwargs)
        raise FileNotFoundError("private executable path")

    monkeypatch.setattr("app.modules.agent1.processing.subprocess.run", missing)
    extractor = FFmpegAudioExtractor("configured-ffmpeg", 17)
    with pytest.raises(VideoProcessorNotConfiguredError):
        extractor.extract(video, output)
    assert captured["arguments"][0] == "configured-ffmpeg"
    assert captured["shell"] is False
    assert captured["timeout"] == 17


def test_ffmpeg_timeout_is_sanitized(tmp_path, monkeypatch):
    video = tmp_path / "video.mov"
    video.write_bytes(b"video")

    def timeout(arguments, **kwargs):
        raise __import__("subprocess").TimeoutExpired(arguments, kwargs["timeout"], stderr=b"private")

    monkeypatch.setattr("app.modules.agent1.processing.subprocess.run", timeout)
    with pytest.raises(VideoAudioExtractionError) as exc_info:
        FFmpegAudioExtractor("ffmpeg", 1).extract(video, tmp_path / "audio.wav")
    assert str(exc_info.value) == "Video audio extraction timed out."
    assert "private" not in str(exc_info.value)


@pytest.mark.parametrize(
    ("error", "code", "message"),
    [
        (VideoProcessorNotConfiguredError("C:\\secret\\ffmpeg"), "VIDEO_PROCESSOR_NOT_CONFIGURED", "Video processing is not configured."),
        (VideoAudioExtractionError("stderr: private path"), "VIDEO_AUDIO_EXTRACTION_FAILED", "Audio could not be extracted from the video."),
        (NoUsableAudioError("private temp path"), "NO_USABLE_AUDIO", "No usable audio transcript could be produced from the video."),
    ],
)
def test_video_extraction_failures_are_safe_and_cleanup_temp_files(
    sources_context, error, code, message
):
    auth, projects, runs, sources, _ = sources_context
    provider = FakeSpeechToTextProvider({"*": "unused transcript"})
    extractor = FakeVideoAudioExtractor(error)
    configure_video_processing(projects, runs, sources, provider, extractor)
    try:
        client = auth.client("USR-001")
        project = create_project(client)
        client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[("files", ("video.webm", b"video", "video/webm"))],
        )
        response = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        )
        result = response.json()["data"][0]
        assert result["processing_status"] == "FAILED"
        assert result["processing_error_code"] == code
        assert result["processing_error_message"] == message
        assert "secret" not in response.text and "private" not in response.text
        assert extractor.output_paths and not extractor.output_paths[0].parent.exists()
        assert runs.documents == {}
    finally:
        app.dependency_overrides.pop(get_requirement_source_processing_service, None)


@pytest.mark.parametrize(
    ("provider_value", "code", "message"),
    [
        (TranscriptionProviderError("provider secret"), "VIDEO_TRANSCRIPTION_FAILED", "The video audio could not be transcribed."),
        ("   \r\n ", "NO_USABLE_AUDIO", "No usable audio transcript could be produced from the video."),
    ],
)
def test_video_transcription_failure_and_blank_transcript_are_safe(
    sources_context, provider_value, code, message
):
    auth, projects, runs, sources, _ = sources_context
    provider = FakeSpeechToTextProvider({"*": provider_value})
    extractor = FakeVideoAudioExtractor()
    configure_video_processing(projects, runs, sources, provider, extractor)
    try:
        client = auth.client("USR-001")
        project = create_project(client)
        client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[("files", ("video.mp4", b"video", "video/mp4"))],
        )
        response = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        )
        result = response.json()["data"][0]
        assert result["processing_status"] == "FAILED"
        assert result["processing_error_code"] == code
        assert result["processing_error_message"] == message
        assert "provider secret" not in response.text
        assert not extractor.output_paths[0].parent.exists()
        assert runs.documents == {}
    finally:
        app.dependency_overrides.pop(get_requirement_source_processing_service, None)


def test_mixed_document_audio_video_partial_failure_without_engine(sources_context):
    auth, projects, runs, sources, _ = sources_context
    provider = FakeSpeechToTextProvider({
        "interview.mp3": "Audio stakeholder requirement",
        "audio.wav": TranscriptionProviderError("video provider secret"),
    })
    extractor = FakeVideoAudioExtractor()
    configure_video_processing(projects, runs, sources, provider, extractor)
    app.dependency_overrides[get_agent1_engine] = lambda: (_ for _ in ()).throw(
        AssertionError("Agent1Engine must not be resolved")
    )
    try:
        client = auth.client("USR-001")
        project = create_project(client)
        client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources",
            files=[
                ("files", ("brief.txt", b"A useful document requirement", "text/plain")),
                ("files", ("interview.mp3", b"audio", "audio/mpeg")),
                ("files", ("walkthrough.mov", b"video", "video/quicktime")),
            ],
        )
        data = client.post(
            f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
        ).json()["data"]
        assert [(item["source_id"], item["processing_status"]) for item in data] == [
            ("DOC-001", "COMPLETED"),
            ("AUD-001", "COMPLETED"),
            ("VID-001", "FAILED"),
        ]
        assert data[2]["processing_error_code"] == "VIDEO_TRANSCRIPTION_FAILED"
        assert sources.documents[(project["agent_project_id"], "DOC-001")]["normalized_text"]
        assert sources.documents[(project["agent_project_id"], "AUD-001")]["normalized_text"]
        assert runs.documents == {}
    finally:
        app.dependency_overrides.pop(get_requirement_source_processing_service, None)
        app.dependency_overrides.pop(get_agent1_engine, None)


def test_processing_enforces_owner_type_and_archived_project(sources_context):
    auth, _, _, _, _ = sources_context
    owner = auth.client("USR-001")
    foreign = auth.client("USR-002")
    project = create_project(owner)
    assert foreign.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
    ).status_code == 404

    wrong = create_project(owner, agent_type="DESIGN_PLANNING")
    assert owner.post(
        f"/api/v1/agents/requirements/{wrong['agent_project_id']}/sources/process"
    ).status_code == 422

    owner.patch(
        f"/api/v1/agent-projects/{project['agent_project_id']}", json={"status": "ARCHIVED"}
    )
    assert owner.post(
        f"/api/v1/agents/requirements/{project['agent_project_id']}/sources/process"
    ).status_code == 409
