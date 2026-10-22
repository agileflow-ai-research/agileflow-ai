import json
from io import BytesIO

import fitz
import pytest
from docx import Document

from app.main import app
from app.modules.agents.implementation.inputs.extraction import constants
from app.modules.agents.implementation.inputs.extraction import service as extraction_service_module
from app.modules.agents.implementation.inputs.extraction.schemas import ExtractionStatus
from app.modules.agents.implementation.api.router import get_implementation_upload_storage
from app.modules.agents.implementation.infrastructure.storage import ImplementationUploadStorage
from app.modules.standalone.router import (
    get_agent_project_repository,
    get_agent_run_repository,
)
from test_standalone_foundation import (
    FakeAgentProjectRepository,
    FakeAgentRunRepository,
)


PNG = b"\x89PNG\r\n\x1a\nminimal"


def pdf_bytes(*pages):
    document = fitz.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def docx_bytes(blocks=()):
    document = Document()
    for block in blocks:
        if block[0] == "heading":
            document.add_heading(block[1], level=block[2])
        elif block[0] == "paragraph":
            document.add_paragraph(block[1])
        else:
            rows = block[1]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture
def extraction_context(authentication_context, tmp_path):
    projects = FakeAgentProjectRepository()
    runs = FakeAgentRunRepository()
    storage = ImplementationUploadStorage(tmp_path)
    app.dependency_overrides[get_agent_project_repository] = lambda: projects
    app.dependency_overrides[get_agent_run_repository] = lambda: runs
    app.dependency_overrides[get_implementation_upload_storage] = lambda: storage
    try:
        yield authentication_context, projects, tmp_path
    finally:
        app.dependency_overrides.pop(get_agent_project_repository, None)
        app.dependency_overrides.pop(get_agent_run_repository, None)
        app.dependency_overrides.pop(get_implementation_upload_storage, None)


def create_project(client, agent_type="IMPLEMENTATION"):
    response = client.post(
        "/api/v1/agent-projects",
        json={"agent_type": agent_type, "name": "Builder", "domain": "SaaS"},
    )
    assert response.status_code == 201
    return response.json()["data"]


def upload(client, project_id, filename, content, source_type="DOCUMENT"):
    response = client.post(
        f"/api/v1/agents/implementation/{project_id}/input-files",
        data={
            "source_type": source_type,
            "artifact_type": "UI_WIREFRAMES" if source_type == "IMAGE" else "API_SPECIFICATION",
        },
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert response.status_code == 201
    return response.json()["data"]


def extract(client, project_id, upload_id, *, force=False):
    return client.post(
        f"/api/v1/agents/implementation/{project_id}/input-files/{upload_id}/extract",
        params={"force": str(force).lower()},
    )


def get_extraction(client, project_id, upload_id):
    return client.get(
        f"/api/v1/agents/implementation/{project_id}/input-files/{upload_id}/extraction"
    )


def upload_dir(root, project_id):
    return root / project_id / "inputs"


def test_extraction_status_enum_is_exact():
    assert {item.value for item in ExtractionStatus} == {"SUCCESS", "PARTIAL", "FAILED"}


def test_pdf_extracts_text_pages_provenance_and_metrics(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    source = pdf_bytes("API Specification", "POST /api/bookings")
    metadata = upload(client, project["agent_project_id"], "quickbook-api.pdf", source)
    response = extract(client, project["agent_project_id"], metadata["upload_id"])
    assert response.status_code == 200
    result = response.json()["data"]
    assert result["status"] == "SUCCESS"
    assert [page["page_number"] for page in result["content"]["pages"]] == [1, 2]
    assert "API Specification" in result["content"]["combined_text"]
    assert result["metrics"] == {
        "page_count": 2,
        "processed_page_count": 2,
        "pages_with_text": 2,
        "pages_without_text": 0,
        "extracted_character_count": len(result["content"]["combined_text"]),
    }
    assert result["source_sha256"] == metadata["sha256"]
    assert result["artifact_type"] == "API_SPECIFICATION"
    assert result["original_filename"] == "quickbook-api.pdf"
    assert "path" not in json.dumps(response.json()).lower()


def test_empty_pdf_fails_and_some_empty_pages_are_partial(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    empty = upload(client, project["agent_project_id"], "empty.pdf", pdf_bytes(""))
    empty_result = extract(client, project["agent_project_id"], empty["upload_id"]).json()["data"]
    assert empty_result["status"] == "FAILED"
    assert {warning["code"] for warning in empty_result["warnings"]} >= {
        "NO_TEXT_ON_PAGE", "NO_EXTRACTABLE_TEXT"
    }

    partial = upload(
        client, project["agent_project_id"], "partial.pdf", pdf_bytes("Useful", "")
    )
    partial_result = extract(client, project["agent_project_id"], partial["upload_id"]).json()["data"]
    assert partial_result["status"] == "PARTIAL"
    assert partial_result["metrics"]["pages_with_text"] == 1
    assert partial_result["metrics"]["pages_without_text"] == 1


def test_pdf_truncates_safely_and_malformed_pdf_fails(extraction_context, monkeypatch):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    monkeypatch.setattr(constants, "MAX_EXTRACTED_CHARACTERS", 12)
    bounded = upload(client, project["agent_project_id"], "bounded.pdf", pdf_bytes("A" * 100))
    result = extract(client, project["agent_project_id"], bounded["upload_id"]).json()["data"]
    assert result["status"] == "PARTIAL"
    assert result["metrics"]["extracted_character_count"] <= 12
    assert any(warning["code"] == "CONTENT_TRUNCATED" for warning in result["warnings"])

    malformed = upload(client, project["agent_project_id"], "broken.pdf", b"%PDF-not-valid")
    failed = extract(client, project["agent_project_id"], malformed["upload_id"]).json()["data"]
    assert failed["status"] == "FAILED"
    assert failed["warnings"][0]["code"] == "EXTRACTION_ERROR"


def test_docx_preserves_heading_paragraph_table_and_order(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    source = docx_bytes(
        [
            ("heading", "API Specification", 1),
            ("paragraph", "Booking endpoints"),
            ("table", [["Method", "Endpoint"], ["POST", "/api/bookings"]]),
            ("paragraph", "Authentication required"),
        ]
    )
    metadata = upload(client, project["agent_project_id"], "spec.docx", source)
    result = extract(client, project["agent_project_id"], metadata["upload_id"]).json()["data"]
    assert result["status"] == "SUCCESS"
    blocks = result["content"]["blocks"]
    assert [block["type"] for block in blocks] == [
        "HEADING", "PARAGRAPH", "TABLE", "PARAGRAPH"
    ]
    assert blocks[0]["level"] == 1
    assert blocks[2]["rows"][1] == ["POST", "/api/bookings"]
    assert [block["block_index"] for block in blocks] == sorted(
        block["block_index"] for block in blocks
    )
    assert result["metrics"]["heading_count"] == 1
    assert result["metrics"]["paragraph_count"] == 2
    assert result["metrics"]["table_count"] == 1


def test_empty_and_truncated_docx(extraction_context, monkeypatch):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    empty = upload(client, project["agent_project_id"], "empty.docx", docx_bytes())
    empty_result = extract(client, project["agent_project_id"], empty["upload_id"]).json()["data"]
    assert empty_result["status"] == "FAILED"
    assert empty_result["warnings"][0]["code"] == "EMPTY_DOCUMENT"

    monkeypatch.setattr(constants, "MAX_EXTRACTED_CHARACTERS", 20)
    large = upload(
        client,
        project["agent_project_id"],
        "large.docx",
        docx_bytes([("paragraph", "content " * 100)]),
    )
    large_result = extract(client, project["agent_project_id"], large["upload_id"]).json()["data"]
    assert large_result["status"] == "PARTIAL"
    assert large_result["metrics"]["extracted_character_count"] <= 20


def test_txt_and_markdown_extraction(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    text = upload(client, project["agent_project_id"], "notes.txt", b"Line one\r\nLine two")
    text_result = extract(client, project["agent_project_id"], text["upload_id"]).json()["data"]
    assert text_result["status"] == "SUCCESS"
    assert text_result["content"]["text"] == "Line one\nLine two"
    assert text_result["metrics"]["line_count"] == 2

    markdown_source = "# API\n\n- item\n\n```js\nfetch('/api')\n```"
    markdown = upload(
        client, project["agent_project_id"], "api.md", markdown_source.encode()
    )
    markdown_result = extract(client, project["agent_project_id"], markdown["upload_id"]).json()["data"]
    assert markdown_result["status"] == "SUCCESS"
    assert markdown_result["content"]["markdown"] == markdown_source
    assert "```js" in markdown_result["content"]["combined_text"]


def test_empty_and_bounded_text(extraction_context, monkeypatch):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    empty = upload(client, project["agent_project_id"], "empty.txt", b"   \n")
    empty_result = extract(client, project["agent_project_id"], empty["upload_id"]).json()["data"]
    assert empty_result["status"] == "FAILED"
    monkeypatch.setattr(constants, "MAX_EXTRACTED_CHARACTERS", 10)
    large = upload(client, project["agent_project_id"], "large.txt", b"x" * 100)
    large_result = extract(client, project["agent_project_id"], large["upload_id"]).json()["data"]
    assert large_result["status"] == "PARTIAL"
    assert large_result["content"]["text"] == "x" * 10


def test_json_is_structured_and_deterministic(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    metadata = upload(
        client,
        project["agent_project_id"],
        "api.json",
        b'{"z": 1, "a": {"method": "POST"}}',
    )
    result = extract(client, project["agent_project_id"], metadata["upload_id"]).json()["data"]
    assert result["status"] == "SUCCESS"
    assert result["content"]["structured_content"] == {
        "z": 1, "a": {"method": "POST"}
    }
    assert result["content"]["normalized_text"] == json.dumps(
        {"z": 1, "a": {"method": "POST"}},
        ensure_ascii=False,
        indent=2,
        sort_keys=True,
    )
    assert result["metrics"]["root_type"] == "OBJECT"


def test_yaml_uses_safe_load_and_returns_json_serializable_content(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    safe = upload(
        client,
        project["agent_project_id"],
        "api.yaml",
        b"defaults: &defaults\n  method: POST\nbooking:\n  <<: *defaults\n  path: /api/bookings\n",
    )
    result = extract(client, project["agent_project_id"], safe["upload_id"]).json()["data"]
    assert result["status"] == "SUCCESS"
    assert result["content"]["structured_content"]["booking"]["method"] == "POST"
    json.dumps(result)

    unsafe = upload(
        client,
        project["agent_project_id"],
        "unsafe.yml",
        b"!!python/object/apply:builtins.str [unsafe]",
    )
    unsafe_result = extract(client, project["agent_project_id"], unsafe["upload_id"]).json()["data"]
    assert unsafe_result["status"] == "FAILED"
    assert unsafe_result["warnings"][0]["code"] == "INVALID_STRUCTURED_CONTENT"
    assert unsafe_result["content"]["structured_content"] is None


def test_valid_cache_hit_avoids_parser_rerun(extraction_context, monkeypatch):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    metadata = upload(client, project["agent_project_id"], "notes.txt", b"cached")
    first = extract(client, project["agent_project_id"], metadata["upload_id"]).json()["data"]

    def should_not_run(*args, **kwargs):
        raise AssertionError("extractor reran")

    monkeypatch.setattr(extraction_service_module, "extract_text", should_not_run)
    second = extract(client, project["agent_project_id"], metadata["upload_id"])
    assert second.status_code == 200
    assert second.json()["data"] == first


@pytest.mark.parametrize("cache_change", ["sha", "version", "malformed"])
def test_invalid_cache_is_ignored_and_rebuilt(
    extraction_context, monkeypatch, cache_change
):
    authentication, _, root = extraction_context
    client = authentication.client()
    project = create_project(client)
    metadata = upload(client, project["agent_project_id"], "notes.txt", b"source")
    extract(client, project["agent_project_id"], metadata["upload_id"])
    sidecar = upload_dir(root, project["agent_project_id"]) / f'{metadata["upload_id"]}.extracted.json'
    if cache_change == "malformed":
        sidecar.write_text("not-json", encoding="utf-8")
    else:
        cached = json.loads(sidecar.read_text(encoding="utf-8"))
        cached["source_sha256" if cache_change == "sha" else "extractor_version"] = (
            "0" * 64 if cache_change == "sha" else "0"
        )
        sidecar.write_text(json.dumps(cached), encoding="utf-8")
    calls = 0
    original = extraction_service_module.extract_text

    def counted(*args, **kwargs):
        nonlocal calls
        calls += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(extraction_service_module, "extract_text", counted)
    rebuilt = extract(client, project["agent_project_id"], metadata["upload_id"])
    assert rebuilt.status_code == 200
    assert calls == 1
    assert rebuilt.json()["data"]["source_sha256"] == metadata["sha256"]
    assert rebuilt.json()["data"]["extractor_version"] == "1"


def test_get_is_read_only_and_returns_only_valid_existing_cache(extraction_context):
    authentication, _, _ = extraction_context
    client = authentication.client()
    project = create_project(client)
    metadata = upload(client, project["agent_project_id"], "notes.txt", b"review")
    assert get_extraction(client, project["agent_project_id"], metadata["upload_id"]).status_code == 404
    expected = extract(client, project["agent_project_id"], metadata["upload_id"]).json()["data"]
    response = get_extraction(client, project["agent_project_id"], metadata["upload_id"])
    assert response.status_code == 200
    assert response.json()["data"] == expected


def test_security_project_and_source_guards(extraction_context):
    authentication, _, _ = extraction_context
    owner = authentication.client("USR-001")
    foreign = authentication.client("USR-002")
    project = create_project(owner)
    metadata = upload(owner, project["agent_project_id"], "notes.txt", b"private")
    assert extract(foreign, project["agent_project_id"], metadata["upload_id"]).status_code == 404
    owner.patch(
        f'/api/v1/agent-projects/{project["agent_project_id"]}',
        json={"status": "ARCHIVED"},
    )
    assert extract(owner, project["agent_project_id"], metadata["upload_id"]).status_code == 409

    wrong = create_project(owner, "DESIGN_PLANNING")
    assert extract(owner, wrong["agent_project_id"], metadata["upload_id"]).status_code == 422

    active = create_project(owner)
    image = upload(owner, active["agent_project_id"], "wireframe.png", PNG, "IMAGE")
    image_response = extract(owner, active["agent_project_id"], image["upload_id"])
    assert image_response.status_code == 422
    assert "Image understanding is not implemented" in image_response.text
    assert extract(
        owner, active["agent_project_id"], "A3UP-0123456789abcdef0123456789abcdef"
    ).status_code == 404
    assert extract(owner, active["agent_project_id"], "../../metadata").status_code == 404


def test_deleting_document_removes_extraction_sidecar(extraction_context):
    authentication, _, root = extraction_context
    client = authentication.client()
    project = create_project(client)
    metadata = upload(client, project["agent_project_id"], "notes.txt", b"delete me")
    extract(client, project["agent_project_id"], metadata["upload_id"])
    directory = upload_dir(root, project["agent_project_id"])
    assert (directory / f'{metadata["upload_id"]}.extracted.json').exists()
    response = client.delete(
        f'/api/v1/agents/implementation/{project["agent_project_id"]}/input-files/{metadata["upload_id"]}'
    )
    assert response.status_code == 200
    assert list(directory.iterdir()) == []
