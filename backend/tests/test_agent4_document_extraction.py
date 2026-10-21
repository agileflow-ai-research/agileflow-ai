import hashlib
import json
from io import BytesIO
from pathlib import Path

import fitz
import pytest
from docx import Document

from app.main import app
from app.modules.agents.testing_validation.preparation.extraction import (
    service as extraction_service_module,
)
from app.modules.agents.testing_validation.preparation.extraction.schemas import (
    Agent4ExtractionStatus,
)
from app.modules.agents.testing_validation.preparation.extraction.service import (
    Agent4DocumentExtractionService,
    Agent4ExtractionGateError,
    Agent4UnsupportedExtractionSourceError,
)
from app.modules.agents.testing_validation.preparation.extraction.storage import (
    Agent4ExtractionStorage,
    Agent4ExtractionStorageError,
)
from app.modules.agents.testing_validation.router import (
    get_agent4_source_storage,
    get_agent4_workflow_storage,
)
from app.modules.agents.testing_validation.sources.storage import (
    Agent4SourceNotFoundError,
    Agent4SourceStorage,
)
from app.modules.agents.testing_validation.workflow.storage import Agent4WorkflowStorage
from app.modules.standalone.router import (
    get_agent_project_repository,
    get_agent_run_repository,
)
from app.modules.agents.testing_validation.workflow.service import (
    Agent4WorkflowStageError,
    WrongAgent4WorkflowProjectError,
)
from app.modules.standalone.service import (
    AgentProjectNotFoundError,
    ArchivedAgentProjectError,
    StandaloneService,
)
from test_standalone_foundation import (
    FakeAgentProjectRepository,
    FakeAgentRunRepository,
    user,
)


def pdf_bytes(*pages: str) -> bytes:
    document = fitz.open()
    for value in pages:
        page = document.new_page()
        if value:
            page.insert_text((72, 72), value)
    result = document.tobytes()
    document.close()
    return result


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Functional Requirements", level=1)
    document.add_paragraph("Customers can transfer funds.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Goal"
    table.cell(1, 0).text = "Customer"
    table.cell(1, 1).text = "Transfer funds"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


@pytest.fixture
def extraction_context(authentication_context, tmp_path):
    projects = FakeAgentProjectRepository()
    runs = FakeAgentRunRepository()
    sources = Agent4SourceStorage(tmp_path)
    workflows = Agent4WorkflowStorage(tmp_path)
    extractions = Agent4ExtractionStorage(tmp_path)
    app.dependency_overrides[get_agent_project_repository] = lambda: projects
    app.dependency_overrides[get_agent_run_repository] = lambda: runs
    app.dependency_overrides[get_agent4_source_storage] = lambda: sources
    app.dependency_overrides[get_agent4_workflow_storage] = lambda: workflows
    service = Agent4DocumentExtractionService(
        StandaloneService(projects, runs), sources, workflows, extractions
    )
    try:
        yield authentication_context.client(), service, tmp_path, projects, runs
    finally:
        app.dependency_overrides.pop(get_agent_project_repository, None)
        app.dependency_overrides.pop(get_agent_run_repository, None)
        app.dependency_overrides.pop(get_agent4_source_storage, None)
        app.dependency_overrides.pop(get_agent4_workflow_storage, None)


def create_project(client) -> str:
    response = client.post(
        "/api/v1/agent-projects",
        json={
            "agent_type": "TESTING_VALIDATION",
            "agent_stage": "TEST_GENERATION",
            "name": "Extraction",
            "domain": "Fintech",
        },
    )
    assert response.status_code == 201
    return response.json()["data"]["agent_project_id"]


def source_url(project_id: str) -> str:
    return f"/api/v1/agents/testing-validation/{project_id}/test-generation/sources"


def upload_document(client, project_id, filename, content, content_type):
    response = client.post(
        f"{source_url(project_id)}/files",
        data={
            "artifact_type": "FUNCTIONAL_REQUIREMENTS",
            "source_type": "DOCUMENT",
            "normally_supplied_by": "AGENT_1",
        },
        files={"file": (filename, content, content_type)},
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def add_text(client, project_id, artifact_type="FUNCTIONAL_REQUIREMENTS"):
    response = client.post(
        f"{source_url(project_id)}/text",
        json={
            "artifact_type": artifact_type,
            "normally_supplied_by": "AGENT_1",
            "content": "Customers can transfer funds.",
        },
    )
    assert response.status_code == 201
    return response.json()["data"]


def confirm(client, project_id):
    selection = client.put(
        f"/api/v1/agents/testing-validation/{project_id}/test-generation/workflow/outputs",
        json={"selected_outputs": ["UNIT_TEST_CASES"]},
    )
    assert selection.status_code == 200
    response = client.post(
        f"/api/v1/agents/testing-validation/{project_id}/test-generation/workflow/n1/confirm"
    )
    assert response.status_code == 200, response.text


def confirmed_document(client, filename, content, content_type):
    project_id = create_project(client)
    source = upload_document(client, project_id, filename, content, content_type)
    confirm(client, project_id)
    return project_id, source


def test_pdf_extracts_pages_and_empty_page_is_partial(extraction_context):
    client, service, root, _, _ = extraction_context
    project_id, source = confirmed_document(
        client, "requirements.pdf", pdf_bytes("First page", "", "Third page"), "application/pdf"
    )
    raw = root / project_id / "inputs" / f"{source['source_id']}.pdf"
    sidecar = raw.with_name(f"{source['source_id']}.source.json")
    workflow = root / project_id / "workflow.json"
    before = (raw.read_bytes(), sidecar.read_bytes(), workflow.read_bytes())

    result = service.extract(project_id, source["source_id"], user())

    assert result.status is Agent4ExtractionStatus.PARTIAL
    assert [page.page_number for page in result.pages] == [1, 2, 3]
    assert result.pages[1].text == ""
    assert {warning.code for warning in result.warnings} >= {"NO_TEXT_ON_PAGE"}
    assert (raw.read_bytes(), sidecar.read_bytes(), workflow.read_bytes()) == before
    persisted = root / project_id / "preparation" / f"{source['source_id']}.extraction.json"
    assert persisted.is_file()
    assert "path" not in result.model_dump_json().lower()


def test_docx_preserves_heading_paragraph_table_and_order(extraction_context):
    client, service, _, _, _ = extraction_context
    project_id, source = confirmed_document(
        client,
        "requirements.docx",
        docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    result = service.extract(project_id, source["source_id"], user())
    assert result.status is Agent4ExtractionStatus.SUCCESS
    assert result.extraction_id.startswith("A4EXT-")
    assert result.agent_project_id == project_id
    assert result.source_id == source["source_id"]
    assert result.artifact_type.value == "FUNCTIONAL_REQUIREMENTS"
    assert result.source_sha256 == source["sha256"]
    assert result.source_size_bytes == source["size_bytes"]
    assert result.extractor_version == "E1"
    assert result.schema_version == "1"
    assert [block.type.value for block in result.blocks] == [
        "HEADING",
        "PARAGRAPH",
        "TABLE",
    ]
    assert result.blocks[0].level == 1
    assert result.blocks[2].rows[1] == ["Customer", "Transfer funds"]
    assert [block.block_index for block in result.blocks] == sorted(
        block.block_index for block in result.blocks
    )


@pytest.mark.parametrize(
    ("filename", "content", "content_type", "expected_format"),
    [
        ("requirements.txt", b"Line one\r\nLine two", "text/plain", "TXT"),
        ("requirements.md", b"# Title\n\n- item\n\n```py\npass\n```", "text/markdown", "MD"),
        ("requirements.json", b'{"z":1,"a":{"value":2}}', "application/json", "JSON"),
        ("requirements.yaml", b"z: 1\na:\n  value: 2\n", "application/yaml", "YAML"),
        ("requirements.yml", b"z: 1\na:\n  value: 2\n", "application/yaml", "YML"),
    ],
)
def test_text_and_structured_formats_extract_deterministically(
    extraction_context, filename, content, content_type, expected_format
):
    client, service, _, _, _ = extraction_context
    project_id, source = confirmed_document(client, filename, content, content_type)
    result = service.extract(project_id, source["source_id"], user())
    assert result.status is Agent4ExtractionStatus.SUCCESS
    assert result.document_format.value == expected_format
    if expected_format in {"TXT", "MD"}:
        assert "Line one\nLine two" in result.extracted_text or "```py" in result.extracted_text
    else:
        assert result.structured_content == {"z": 1, "a": {"value": 2}}
        assert result.extracted_text == json.dumps(
            result.structured_content,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )


@pytest.mark.parametrize(
    ("filename", "content", "content_type", "error_code"),
    [
        ("broken.pdf", b"%PDF-not-valid", "application/pdf", "INVALID_PDF"),
        (
            "unsafe.yaml",
            b"!!python/object/apply:builtins.str [unsafe]",
            "application/yaml",
            "INVALID_YAML",
        ),
        ("bad.yaml", b"key: [unterminated", "application/yaml", "INVALID_YAML"),
    ],
)
def test_malformed_documents_fail_safely(
    extraction_context, filename, content, content_type, error_code
):
    client, service, _, _, _ = extraction_context
    project_id, source = confirmed_document(client, filename, content, content_type)
    result = service.extract(project_id, source["source_id"], user())
    assert result.status is Agent4ExtractionStatus.FAILED
    assert result.error.code.value == error_code
    assert "path" not in result.model_dump_json().lower()


def test_docx_archive_limit_and_text_encoding_fail_safely(
    extraction_context, monkeypatch
):
    client, service, _, _, _ = extraction_context
    project_id, source = confirmed_document(
        client,
        "requirements.docx",
        docx_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    monkeypatch.setattr(extraction_service_module, "MAX_DOCX_ARCHIVE_ENTRIES", 1)
    result = service.extract(project_id, source["source_id"], user(), force=True)
    assert result.status is Agent4ExtractionStatus.FAILED
    assert result.error.code.value == "INVALID_DOCX"

    project_id, source = confirmed_document(
        client, "requirements.txt", b"valid", "text/plain"
    )
    raw = Path(service._sources.locate(
        agent_project_id=project_id, owner_user_id="USR-001", source_id=source["source_id"]
    ).file_path)
    bad = b"\xff\xfe"
    raw.write_bytes(bad)
    metadata_path = raw.with_name(f"{source['source_id']}.source.json")
    payload = json.loads(metadata_path.read_text(encoding="utf-8"))
    payload["size_bytes"] = len(bad)
    payload["sha256"] = hashlib.sha256(bad).hexdigest()
    metadata_path.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
    workflow_path = raw.parents[1] / "workflow.json"
    workflow_payload = json.loads(workflow_path.read_text(encoding="utf-8"))
    from app.modules.agents.testing_validation.sources.schemas import TestGenerationSourceRecord
    from app.modules.agents.testing_validation.workflow.digests import source_set_digest
    workflow_payload["n1_confirmation"]["source_set_digest"] = source_set_digest(
        [
            TestGenerationSourceRecord.model_validate(
                {k: v for k, v in payload.items() if k != "owner_user_id"}
            )
        ]
    )
    workflow_path.write_text(json.dumps(workflow_payload, separators=(",", ":")), encoding="utf-8")
    result = service.extract(project_id, source["source_id"], user(), force=True)
    assert result.error.code.value == "ENCODING_ERROR"


def test_confirmation_and_snapshot_membership_are_required(extraction_context):
    client, service, _, _, _ = extraction_context
    project_id = create_project(client)
    source = upload_document(client, project_id, "requirements.txt", b"text", "text/plain")
    with pytest.raises(Agent4ExtractionGateError):
        service.extract(project_id, source["source_id"], user())
    confirm(client, project_id)
    with pytest.raises(Agent4SourceNotFoundError):
        service.extract(project_id, "A4SRC-" + "f" * 32, user())

    other_project, other_source = confirmed_document(
        client, "other.txt", b"other", "text/plain"
    )
    with pytest.raises(Agent4SourceNotFoundError):
        service.extract(project_id, other_source["source_id"], user())
    assert other_project != project_id


def test_confirmed_source_metadata_digest_must_still_match(extraction_context):
    client, service, root, _, _ = extraction_context
    project_id, source = confirmed_document(
        client, "requirements.txt", b"approved", "text/plain"
    )
    sidecar = root / project_id / "inputs" / f"{source['source_id']}.source.json"
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    payload["artifact_type"] = "SRS"
    sidecar.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
    with pytest.raises(Agent4ExtractionGateError, match="confirmed setup"):
        service.extract(project_id, source["source_id"], user())


def test_project_owner_lifecycle_type_and_stage_boundaries(extraction_context):
    client, service, _, projects, _ = extraction_context
    project_id, source = confirmed_document(
        client, "requirements.txt", b"private", "text/plain"
    )
    with pytest.raises(AgentProjectNotFoundError):
        service.extract(project_id, source["source_id"], user("USR-002"))

    projects.documents[project_id]["status"] = "ARCHIVED"
    with pytest.raises(ArchivedAgentProjectError):
        service.extract(project_id, source["source_id"], user())
    projects.documents[project_id]["status"] = "ACTIVE"
    projects.documents[project_id]["agent_stage"] = "VALIDATION"
    with pytest.raises(Agent4WorkflowStageError):
        service.extract(project_id, source["source_id"], user())
    projects.documents[project_id]["agent_stage"] = None
    with pytest.raises(Agent4WorkflowStageError):
        service.extract(project_id, source["source_id"], user())

    with pytest.raises(Agent4ExtractionGateError):
        service.extract("not-an-id", source["source_id"], user())

    response = client.post(
        "/api/v1/agent-projects",
        json={"agent_type": "IMPLEMENTATION", "name": "Wrong", "domain": "SaaS"},
    )
    wrong_project = response.json()["data"]["agent_project_id"]
    with pytest.raises(WrongAgent4WorkflowProjectError):
        service.extract(wrong_project, source["source_id"], user())


def test_missing_size_and_hash_mismatch_are_structured_failures(extraction_context):
    client, service, root, _, _ = extraction_context
    project_id, source = confirmed_document(client, "requirements.txt", b"evidence", "text/plain")
    raw = root / project_id / "inputs" / f"{source['source_id']}.txt"
    raw.unlink()
    missing = service.extract(project_id, source["source_id"], user())
    assert missing.error.code.value == "SOURCE_FILE_MISSING"

    project_id, source = confirmed_document(client, "requirements.txt", b"evidence", "text/plain")
    raw = root / project_id / "inputs" / f"{source['source_id']}.txt"
    raw.write_bytes(b"longer evidence")
    size = service.extract(project_id, source["source_id"], user())
    assert size.error.code.value == "SOURCE_SIZE_MISMATCH"

    project_id, source = confirmed_document(client, "requirements.txt", b"evidence", "text/plain")
    raw = root / project_id / "inputs" / f"{source['source_id']}.txt"
    raw.write_bytes(b"tampered")
    digest = service.extract(project_id, source["source_id"], user())
    assert digest.error.code.value == "SOURCE_HASH_MISMATCH"


def test_cache_staleness_and_force_rebuild(extraction_context, monkeypatch):
    client, service, root, _, _ = extraction_context
    project_id, source = confirmed_document(client, "requirements.txt", b"cached", "text/plain")
    first = service.extract(project_id, source["source_id"], user())

    def should_not_run(*args, **kwargs):
        raise AssertionError("parser reran")

    original = extraction_service_module.extract_text
    monkeypatch.setattr(extraction_service_module, "extract_text", should_not_run)
    cached = service.extract(project_id, source["source_id"], user())
    assert cached == first
    assert cached.extracted_at == first.extracted_at

    monkeypatch.setattr(extraction_service_module, "extract_text", original)
    forced = service.extract(project_id, source["source_id"], user(), force=True)
    assert forced.extraction_id != first.extraction_id
    sidecar = root / project_id / "preparation" / f"{source['source_id']}.extraction.json"
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    payload["extractor_version"] = "stale"
    sidecar.write_text(json.dumps(payload), encoding="utf-8")
    rebuilt = service.extract(project_id, source["source_id"], user())
    assert rebuilt.extraction_id != forced.extraction_id

    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    payload["source_sha256"] = "0" * 64
    sidecar.write_text(json.dumps(payload), encoding="utf-8")
    rebuilt_again = service.extract(project_id, source["source_id"], user())
    assert rebuilt_again.extraction_id != rebuilt.extraction_id


def test_usable_truncated_document_is_explicitly_partial(
    extraction_context, monkeypatch
):
    client, service, _, _, _ = extraction_context
    project_id, source = confirmed_document(
        client, "requirements.txt", b"meaningful content", "text/plain"
    )
    monkeypatch.setattr(extraction_service_module, "MAX_EXTRACTED_CHARACTERS", 5)
    result = service.extract(project_id, source["source_id"], user(), force=True)
    assert result.status is Agent4ExtractionStatus.PARTIAL
    assert result.extracted_text == "meani"
    assert result.truncated is True
    assert [warning.code for warning in result.warnings] == ["CONTENT_TRUNCATED"]


def test_text_and_image_sources_create_no_extraction_sidecar(extraction_context):
    client, service, root, _, _ = extraction_context
    project_id = create_project(client)
    text_source = add_text(client, project_id)
    image_response = client.post(
        f"{source_url(project_id)}/files",
        data={
            "artifact_type": "WIREFRAMES",
            "source_type": "IMAGE",
            "normally_supplied_by": "AGENT_2",
        },
        files={"file": ("wireframe.png", b"\x89PNG\r\n\x1a\nminimal", "image/png")},
    )
    assert image_response.status_code == 201
    image_source = image_response.json()["data"]
    confirm(client, project_id)
    with pytest.raises(Agent4UnsupportedExtractionSourceError):
        service.extract(project_id, text_source["source_id"], user())
    with pytest.raises(Agent4UnsupportedExtractionSourceError):
        service.extract(project_id, image_source["source_id"], user())
    assert image_source["collection_status"] == "PENDING_ANALYSIS"
    assert not (root / project_id / "preparation").exists()


def test_atomic_failure_cleans_partial_and_workflow_run_state_is_unchanged(
    extraction_context, monkeypatch
):
    client, service, root, _, runs = extraction_context
    project_id, source = confirmed_document(client, "requirements.txt", b"atomic", "text/plain")
    workflow = root / project_id / "workflow.json"
    workflow_before = workflow.read_bytes()
    original_replace = Path.replace

    def fail_replace(path, target):
        if str(path).endswith(".extraction.json.part"):
            raise OSError("injected")
        return original_replace(path, target)

    monkeypatch.setattr(Path, "replace", fail_replace)
    with pytest.raises(Agent4ExtractionStorageError):
        service.extract(project_id, source["source_id"], user())
    preparation = root / project_id / "preparation"
    assert not list(preparation.glob("*.part"))
    assert workflow.read_bytes() == workflow_before
    state = json.loads(workflow_before)
    assert state["current_node_id"] == "INPUT_PREPARATION"
    assert state["nodes"][1]["status"] == "NOT_STARTED"
    assert runs.documents == {}
    assert not list((root / project_id).rglob("*normal*"))
    assert not list((root / project_id).rglob("*chunk*"))


def test_agent4_router_has_no_extraction_execution_route():
    paths = {
        route.path
        for route in app.routes
        if hasattr(route, "path") and "testing-validation" in route.path
    }
    assert not any("extract" in path or "prepare" in path for path in paths)
