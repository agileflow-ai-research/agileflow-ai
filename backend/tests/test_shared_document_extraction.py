"""Tests for the shared document extraction primitive, independent of any agent."""

import pathlib
from io import BytesIO

import fitz
import pytest
from docx import Document
from pydantic import ValidationError

from app.common.document_extraction import dispatch
from app.common.document_extraction.constants import DEFAULT_LIMITS, ExtractionLimits
from app.common.document_extraction.docx_extractor import extract_docx
from app.common.document_extraction.pdf_extractor import extract_pdf
from app.common.document_extraction.schemas import (
    DocumentType,
    DOCXContent,
    ExtractionStatus,
    ExtractionWarningCode,
    MarkdownContent,
    PDFContent,
    StructuredContent,
    TextContent,
    TextMetrics,
    validate_content_matches_type,
)
from app.common.document_extraction.structured_extractors import extract_structured, extract_text


def write(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return path


def pdf_bytes(*pages):
    document = fitz.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def docx_bytes():
    document = Document()
    document.add_heading("Checkout", level=1)
    document.add_paragraph("Customers pay with a saved card.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Rule"
    table.cell(1, 0).text = "Amount"
    table.cell(1, 1).text = "Positive"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def codes(warnings):
    return {warning.code for warning in warnings}


def test_pdf_extracts_pages_and_text(tmp_path):
    path = write(tmp_path, "a.pdf", pdf_bytes("Booking requires a code"))
    status, content, metrics, warnings = extract_pdf(path)
    assert status is ExtractionStatus.SUCCESS
    assert isinstance(content, PDFContent)
    assert "Booking requires a code" in content.combined_text
    assert metrics.page_count == 1 and metrics.pages_with_text == 1 and warnings == []


def test_pdf_with_some_empty_pages_is_partial_not_success(tmp_path):
    path = write(tmp_path, "mixed.pdf", pdf_bytes("Has text", None))
    status, _, metrics, warnings = extract_pdf(path)
    assert status is ExtractionStatus.PARTIAL
    assert metrics.pages_without_text == 1
    assert ExtractionWarningCode.NO_TEXT_ON_PAGE in codes(warnings)


def test_pdf_without_any_text_fails_with_warning(tmp_path):
    path = write(tmp_path, "blank.pdf", pdf_bytes(None))
    status, _, _, warnings = extract_pdf(path)
    assert status is ExtractionStatus.FAILED
    assert ExtractionWarningCode.NO_EXTRACTABLE_TEXT in codes(warnings)


def test_pdf_page_and_character_limits_are_applied_through_arguments(tmp_path):
    path = write(tmp_path, "two.pdf", pdf_bytes("First page text", "Second page text"))
    status, _, metrics, warnings = extract_pdf(path, limits=ExtractionLimits(max_pdf_pages=1))
    assert status is ExtractionStatus.PARTIAL
    assert metrics.processed_page_count == 1
    assert ExtractionWarningCode.CONTENT_TRUNCATED in codes(warnings)

    status, content, _, warnings = extract_pdf(path, limits=ExtractionLimits(max_extracted_characters=8))
    assert status is ExtractionStatus.PARTIAL
    assert len(content.combined_text) <= 8
    assert ExtractionWarningCode.CONTENT_TRUNCATED in codes(warnings)


def test_malformed_pdf_is_a_safe_failed_outcome(tmp_path):
    path = write(tmp_path, "broken.pdf", b"%PDF-1.7 not really a pdf")
    status, content, metrics, warnings = dispatch.extract_document_safely(path, DocumentType.PDF)
    assert status is ExtractionStatus.FAILED
    assert content.combined_text == ""
    assert metrics.extracted_character_count == 0
    assert ExtractionWarningCode.EXTRACTION_ERROR in codes(warnings)


def test_docx_preserves_heading_paragraph_and_table_in_order(tmp_path):
    path = write(tmp_path, "a.docx", docx_bytes())
    status, content, metrics, _ = extract_docx(path)
    assert status is ExtractionStatus.SUCCESS
    assert isinstance(content, DOCXContent)
    assert [block.type.value for block in content.blocks] == ["HEADING", "PARAGRAPH", "TABLE"]
    assert metrics.table_count == 1 and metrics.heading_count == 1


def test_docx_archive_entry_limit_raises_and_safe_dispatch_fails(tmp_path):
    path = write(tmp_path, "a.docx", docx_bytes())
    with pytest.raises(ValueError):
        extract_docx(path, limits=ExtractionLimits(max_docx_archive_entries=1))
    status, _, _, warnings = dispatch.extract_document_safely(
        path, DocumentType.DOCX, limits=ExtractionLimits(max_docx_archive_entries=1)
    )
    assert status is ExtractionStatus.FAILED
    assert ExtractionWarningCode.EXTRACTION_ERROR in codes(warnings)


def test_docx_uncompressed_size_limit_is_enforced(tmp_path):
    path = write(tmp_path, "a.docx", docx_bytes())
    with pytest.raises(ValueError):
        extract_docx(path, limits=ExtractionLimits(max_docx_uncompressed_bytes=1))


def test_empty_docx_fails(tmp_path):
    document = Document()
    output = BytesIO()
    document.save(output)
    path = write(tmp_path, "empty.docx", output.getvalue())
    status, _, _, warnings = extract_docx(path)
    assert status is ExtractionStatus.FAILED
    assert ExtractionWarningCode.EMPTY_DOCUMENT in codes(warnings)


def test_txt_and_markdown_keep_text_and_use_their_content_types(tmp_path):
    plain = write(tmp_path, "a.txt", "Refund within 14 days\r\nSecond line")
    status, content, metrics, _ = extract_text(plain)
    assert status is ExtractionStatus.SUCCESS
    assert isinstance(content, TextContent)
    assert content.text == "Refund within 14 days\nSecond line"
    assert metrics.line_count == 2

    markdown = write(tmp_path, "a.md", "# Title\n\n- item")
    status, content, _, _ = extract_text(markdown, markdown=True)
    assert status is ExtractionStatus.SUCCESS
    assert isinstance(content, MarkdownContent)
    assert content.markdown.startswith("# Title")


def test_empty_and_bounded_text(tmp_path):
    empty = write(tmp_path, "empty.txt", "   \n  ")
    status, _, _, warnings = extract_text(empty)
    assert status is ExtractionStatus.FAILED
    assert ExtractionWarningCode.EMPTY_DOCUMENT in codes(warnings)

    long_text = write(tmp_path, "long.txt", "x" * 50)
    status, content, metrics, warnings = extract_text(long_text, limits=ExtractionLimits(max_extracted_characters=10))
    assert status is ExtractionStatus.PARTIAL
    assert len(content.text) == 10 and metrics.extracted_character_count == 10
    assert ExtractionWarningCode.CONTENT_TRUNCATED in codes(warnings)


def test_json_is_deterministic_and_malformed_json_fails_safely(tmp_path):
    source = '{"b": 2, "a": {"z": 1, "y": [1, 2]}}'
    path = write(tmp_path, "a.json", source)
    first = extract_structured(path)
    second = extract_structured(path)
    assert first[0] is ExtractionStatus.SUCCESS
    assert first[1].normalized_text == second[1].normalized_text
    assert first[1].normalized_text.index('"a"') < first[1].normalized_text.index('"b"')

    broken = write(tmp_path, "broken.json", '{"a": ')
    status, content, _, warnings = extract_structured(broken)
    assert status is ExtractionStatus.FAILED
    assert content.normalized_text == ""
    assert ExtractionWarningCode.INVALID_STRUCTURED_CONTENT in codes(warnings)


def test_yaml_uses_safe_load_and_rejects_executable_tags(tmp_path):
    path = write(tmp_path, "a.yaml", "paths:\n  /bookings:\n    get: {}\n")
    status, content, metrics, _ = extract_structured(path, yaml_document=True)
    assert status is ExtractionStatus.SUCCESS
    assert content.structured_content == {"paths": {"/bookings": {"get": {}}}}
    assert metrics.root_type == "OBJECT"

    unsafe = write(tmp_path, "unsafe.yaml", "!!python/object/apply:os.system ['echo unsafe']\n")
    status, _, _, warnings = extract_structured(unsafe, yaml_document=True)
    assert status is ExtractionStatus.FAILED
    assert ExtractionWarningCode.INVALID_STRUCTURED_CONTENT in codes(warnings)


def test_dispatch_routes_every_document_type_to_its_content_model(tmp_path):
    files = {
        DocumentType.PDF: write(tmp_path, "a.pdf", pdf_bytes("Text")),
        DocumentType.DOCX: write(tmp_path, "a.docx", docx_bytes()),
        DocumentType.TXT: write(tmp_path, "a.txt", "Text"),
        DocumentType.MD: write(tmp_path, "a.md", "# Text"),
        DocumentType.JSON: write(tmp_path, "a.json", '{"a": 1}'),
        DocumentType.YAML: write(tmp_path, "a.yaml", "a: 1\n"),
    }
    for document_type, path in files.items():
        status, content, metrics, _ = dispatch.extract_document(path, document_type)
        validate_content_matches_type(document_type, content, metrics)
        assert status is ExtractionStatus.SUCCESS


def test_extension_mapping_covers_exactly_the_contract_formats():
    assert set(dispatch.DOCUMENT_TYPES) == {".pdf", ".docx", ".txt", ".md", ".json", ".yaml", ".yml"}
    assert dispatch.document_type_for_extension(".png") is None


def test_content_shape_validation_rejects_mismatched_models():
    with pytest.raises(ValueError):
        validate_content_matches_type(DocumentType.TXT, StructuredContent(structured_content=None, normalized_text=""), TextMetrics(extracted_character_count=0, line_count=0))


def test_failed_outcome_is_safe_and_typed_for_every_format():
    for document_type in DocumentType:
        status, content, metrics, warnings = dispatch.failed_outcome(document_type)
        validate_content_matches_type(document_type, content, metrics)
        assert status is ExtractionStatus.FAILED
        assert warnings[0].code is ExtractionWarningCode.EXTRACTION_ERROR


def test_default_limits_match_the_agent_3_proven_values():
    assert DEFAULT_LIMITS == ExtractionLimits(
        max_pdf_pages=500,
        max_extracted_characters=2_000_000,
        max_docx_uncompressed_bytes=50 * 1024 * 1024,
        max_docx_archive_entries=10_000,
    )


def test_shared_package_has_no_agent_or_persistence_coupling():
    package = pathlib.Path(__file__).resolve().parents[1] / "app" / "common" / "document_extraction"
    forbidden = ("modules.agents", "modules.standalone", "modules.auth", "A3UP", "A4UP", "AgentProject")
    for path in package.glob("*.py"):
        text = path.read_text(encoding="utf-8")
        for token in forbidden:
            assert token not in text, f"{path.name} must not reference {token}"
