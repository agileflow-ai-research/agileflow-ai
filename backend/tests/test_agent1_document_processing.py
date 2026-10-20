from io import BytesIO

import fitz
import pytest
from docx import Document

from app.main import app
from app.modules.agents.requirements.api.dependencies import (
    get_requirement_source_storage,
)
from app.modules.agents.requirements.infrastructure.storage import (
    RequirementSourceStorage,
)
from app.modules.agents.requirements.processing import service as processing_module
from app.modules.standalone.router import (
    get_agent_project_repository,
    get_agent_run_repository,
)
from test_standalone_foundation import (
    FakeAgentProjectRepository,
    FakeAgentRunRepository,
)


def pdf_bytes(text: str | None = None) -> bytes:
    document = fitz.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    value = document.tobytes()
    document.close()
    return value


def docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("Patients can book appointments.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Goal"
    table.cell(1, 0).text = "Patient"
    table.cell(1, 1).text = "Book online"
    document.save(output)
    return output.getvalue()


@pytest.fixture
def processing_context(authentication_context, tmp_path):
    projects = FakeAgentProjectRepository()
    runs = FakeAgentRunRepository()
    storage = RequirementSourceStorage(tmp_path)
    app.dependency_overrides[get_agent_project_repository] = lambda: projects
    app.dependency_overrides[get_agent_run_repository] = lambda: runs
    app.dependency_overrides[get_requirement_source_storage] = lambda: storage
    try:
        yield authentication_context, projects, runs, tmp_path
    finally:
        app.dependency_overrides.pop(get_agent_project_repository, None)
        app.dependency_overrides.pop(get_agent_run_repository, None)
        app.dependency_overrides.pop(get_requirement_source_storage, None)


def create_project(client, agent_type="REQUIREMENTS"):
    response = client.post(
        "/api/v1/agent-projects",
        json={"agent_type": agent_type, "name": "Discovery", "domain": "Health"},
    )
    assert response.status_code == 201
    return response.json()["data"]


def upload(client, project_id, filename, content, source_type="DOCUMENT"):
    return client.post(
        f"/api/v1/agents/requirements/{project_id}/sources",
        data={"source_type": source_type},
        files={"file": (filename, content, "application/octet-stream")},
    )


def process(client, project_id, source_id, *, force=False):
    suffix = "?force=true" if force else ""
    return client.post(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/processing{suffix}"
    )


def status(client, project_id, source_id):
    return client.get(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/processing/status"
    )


def result(client, project_id, source_id):
    return client.get(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/processing/result"
    )


def test_txt_extraction_and_ready_completed_status(processing_context):
    authentication, _, runs, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "requirements.txt", b"Line one\nLine two").json()["data"]

    ready = status(client, project_id, source["source_id"])
    completed = process(client, project_id, source["source_id"])
    extracted = result(client, project_id, source["source_id"])

    assert ready.json()["data"]["status"] == "READY"
    assert completed.json()["data"]["status"] == "COMPLETED"
    assert extracted.json()["data"]["text"] == "Line one\nLine two"
    assert extracted.json()["data"]["document_type"] == "TXT"
    assert extracted.json()["data"]["source_id"] == "DOC-001"
    assert runs.documents == {}


@pytest.mark.parametrize(
    "content,expected",
    [
        (b"\xef\xbb\xbfUTF-8 BOM text", "UTF-8 BOM text"),
        ("UTF-16 BOM text".encode("utf-16"), "UTF-16 BOM text"),
    ],
)
def test_text_bom_decoding(processing_context, content, expected):
    authentication, _, _, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "bom.txt", content).json()["data"]

    assert process(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"
    assert result(client, project_id, source["source_id"]).json()["data"]["text"] == expected


def test_pdf_machine_readable_extraction(processing_context):
    authentication, _, _, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "brief.pdf", pdf_bytes("Machine readable requirement")).json()["data"]

    assert process(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"
    extracted = result(client, project_id, source["source_id"]).json()["data"]
    assert extracted["text"] == "Machine readable requirement"
    assert extracted["page_count"] == 1


def test_docx_paragraphs_and_tables(processing_context):
    authentication, _, _, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "requirements.docx", docx_bytes()).json()["data"]

    process(client, project_id, source["source_id"])
    extracted = result(client, project_id, source["source_id"]).json()["data"]

    assert "Patients can book appointments." in extracted["text"]
    assert "Role\tGoal" in extracted["text"]
    assert "Patient\tBook online" in extracted["text"]
    assert extracted["paragraph_count"] == 1
    assert extracted["table_count"] == 1


def test_scanned_or_empty_pdf_fails_without_ocr(processing_context):
    authentication, _, _, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "scan.pdf", pdf_bytes()).json()["data"]

    response = process(client, project_id, source["source_id"])

    assert response.status_code == 200
    record = response.json()["data"]
    assert record["status"] == "FAILED"
    assert record["error"]["code"] == "NO_MACHINE_READABLE_TEXT"
    assert "OCR" in record["error"]["message"]
    assert result(client, project_id, source["source_id"]).status_code == 404


def test_unreadable_document_returns_safe_failure(processing_context):
    authentication, _, _, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "broken.pdf", b"%PDF-1.7\nnot-a-real-document").json()["data"]

    record = process(client, project_id, source["source_id"]).json()["data"]

    assert record["status"] == "FAILED"
    assert record["error"] == {
        "code": "DOCUMENT_UNREADABLE",
        "message": "The document could not be read safely.",
    }
    assert "path" not in str(record).lower()


def test_retry_reprocesses_failed_document(processing_context, monkeypatch):
    authentication, _, _, _ = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "retry.txt", b"Retry content").json()["data"]
    original = processing_module.extract_text
    monkeypatch.setattr(
        processing_module,
        "extract_text",
        lambda path: (_ for _ in ()).throw(ValueError("parser unavailable")),
    )
    assert process(client, project_id, source["source_id"]).json()["data"]["status"] == "FAILED"
    monkeypatch.setattr(processing_module, "extract_text", original)

    retried = process(client, project_id, source["source_id"], force=True)

    assert retried.json()["data"]["status"] == "COMPLETED"
    assert result(client, project_id, source["source_id"]).json()["data"]["text"] == "Retry content"


def test_checksum_and_processor_version_bind_cached_results(processing_context, monkeypatch):
    authentication, _, _, root = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "bound.txt", b"Bound content").json()["data"]
    process(client, project_id, source["source_id"])
    directory = root / project_id / "inputs"
    processing_path = directory / "DOC-001.document-processing.json"
    result_path = directory / "DOC-001.document-result.json"

    cached_result = result_path.read_text(encoding="utf-8")
    result_path.write_text(
        cached_result.replace(source["sha256"], "0" * 64), encoding="utf-8"
    )
    assert result(client, project_id, source["source_id"]).status_code == 404
    assert process(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"

    cached_status = processing_path.read_text(encoding="utf-8")
    processing_path.write_text(
        cached_status.replace('"processor_version":"1"', '"processor_version":"999"'),
        encoding="utf-8",
    )
    assert status(client, project_id, source["source_id"]).json()["data"]["status"] == "READY"


def test_changed_source_bytes_fail_checksum_without_mutating_metadata(processing_context):
    authentication, _, runs, root = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "immutable.txt", b"Original").json()["data"]
    metadata_path = root / project_id / "inputs" / "DOC-001.metadata.json"
    original_metadata = metadata_path.read_bytes()
    (root / project_id / "inputs" / "DOC-001.txt").write_bytes(b"Changed")

    record = process(client, project_id, source["source_id"]).json()["data"]

    assert record["status"] == "FAILED"
    assert record["error"]["code"] == "SOURCE_CHECKSUM_MISMATCH"
    assert metadata_path.read_bytes() == original_metadata
    assert runs.documents == {}


def test_owner_type_archive_and_document_guards(processing_context):
    authentication, _, _, _ = processing_context
    owner = authentication.client("USR-001")
    foreign = authentication.client("USR-002")
    project_id = create_project(owner)["agent_project_id"]
    document = upload(owner, project_id, "requirements.txt", b"Requirement").json()["data"]
    video = upload(owner, project_id, "workflow.mp4", b"\x00\x00\x00\x18ftypmp42video", "VIDEO").json()["data"]

    assert process(foreign, project_id, document["source_id"]).status_code == 404
    assert process(owner, project_id, video["source_id"]).status_code == 422
    wrong_id = create_project(owner, "IMPLEMENTATION")["agent_project_id"]
    assert status(owner, wrong_id, "DOC-001").status_code == 422

    owner.patch(f"/api/v1/agent-projects/{project_id}", json={"status": "ARCHIVED"})
    assert process(owner, project_id, document["source_id"]).status_code == 409
    assert status(owner, project_id, document["source_id"]).status_code == 409
    assert result(owner, project_id, document["source_id"]).status_code == 409


def test_source_metadata_is_unchanged_and_processing_survives_reload(processing_context):
    authentication, _, _, root = processing_context
    client = authentication.client()
    project_id = create_project(client)["agent_project_id"]
    source = upload(client, project_id, "persisted.txt", b"Persisted text").json()["data"]
    listed_before = client.get(f"/api/v1/agents/requirements/{project_id}/sources").json()["data"]
    process(client, project_id, source["source_id"])
    listed_after = client.get(f"/api/v1/agents/requirements/{project_id}/sources").json()["data"]

    app.dependency_overrides[get_requirement_source_storage] = lambda: RequirementSourceStorage(root)

    assert listed_after == listed_before
    assert status(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"
    assert result(client, project_id, source["source_id"]).json()["data"]["text"] == "Persisted text"
