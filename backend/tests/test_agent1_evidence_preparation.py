import json
import sys
from datetime import datetime, timezone
from importlib.metadata import PackageNotFoundError
from io import BytesIO
from types import SimpleNamespace

import fitz
import pytest
from docx import Document

from app.main import app
from app.modules.agents.requirements.api.dependencies import (
    get_requirement_source_storage,
    get_sentence_segmenter,
    get_speech_to_text_provider,
    get_video_audio_extractor,
)
from app.modules.agents.requirements.infrastructure.storage import (
    RequirementSourceStorage,
)
from app.modules.agents.requirements.contracts.schemas import RequirementSourceType
from app.modules.agents.requirements.processing.audio.providers.base import (
    SpeechToTextConfiguration,
    SpeechToTextResult,
    SpeechToTextSegment,
)
from app.modules.agents.requirements.processing.preparation.chunker import (
    DeterministicChunker,
    SentenceSpan,
    SentenceSegmenter,
    build_chunker_configuration,
)
from app.modules.agents.requirements.processing.preparation import chunker as chunker_module
from app.modules.agents.requirements.processing.preparation.normalizer import (
    normalize_text,
)
from app.modules.agents.requirements.processing.preparation.schemas import (
    NormalizedEvidenceResult,
    NormalizedProvenanceSpan,
)
from app.modules.agents.requirements.processing.video.providers.base import (
    VideoAudioExtractionConfiguration,
)
from app.modules.standalone.router import (
    get_agent_project_repository,
    get_agent_run_repository,
)
from test_standalone_foundation import (
    FakeAgentProjectRepository,
    FakeAgentRunRepository,
)


class FakeSpeechToTextProvider:
    def __init__(self) -> None:
        self.calls = 0
        self.configuration = SpeechToTextConfiguration(
            provider="fake",
            provider_version="test-1",
            model_name="fake-small",
            device="cpu",
            compute_type="int8",
            language=None,
            fingerprint="a" * 64,
        )

    def transcribe(self, audio_path):
        assert audio_path.is_file()
        self.calls += 1
        return SpeechToTextResult(
            transcript=" First stakeholder request. Second constraint.",
            segments=(
                SpeechToTextSegment(0.0, 1.5, " First stakeholder request."),
                SpeechToTextSegment(1.5, 3.0, " Second constraint."),
            ),
            detected_language="en",
            language_probability=0.99,
        )


class FakeVideoAudioExtractor:
    def __init__(self) -> None:
        self.calls = 0
        self.configuration = VideoAudioExtractionConfiguration(
            provider="fake_ffmpeg",
            provider_version="test-1",
            executable="fake-ffmpeg",
            timeout_seconds=10,
            sample_rate_hz=16_000,
            channels=1,
            fingerprint="b" * 64,
        )

    def extract_audio(self, input_video_path, output_audio_path):
        assert input_video_path.is_file()
        self.calls += 1
        output_audio_path.write_bytes(b"RIFFfake-wave")


class CheckingSentenceSegmenter(SentenceSegmenter):
    def __init__(self, root) -> None:
        super().__init__()
        self.root = root
        self.calls = 0

    @property
    def identity(self):
        return "checking_sentencizer", "test-1"

    def segment(self, text):
        records = [
            json.loads(path.read_text(encoding="utf-8"))
            for path in self.root.glob("*/inputs/*.evidence-preparation.json")
        ]
        assert any(record["status"] == "PROCESSING" for record in records)
        self.calls += 1
        return super().segment(text)


@pytest.fixture
def preparation_context(authentication_context, tmp_path):
    projects = FakeAgentProjectRepository()
    runs = FakeAgentRunRepository()
    storage = RequirementSourceStorage(tmp_path)
    transcription = FakeSpeechToTextProvider()
    extractor = FakeVideoAudioExtractor()
    segmenter = CheckingSentenceSegmenter(tmp_path)
    app.dependency_overrides[get_agent_project_repository] = lambda: projects
    app.dependency_overrides[get_agent_run_repository] = lambda: runs
    app.dependency_overrides[get_requirement_source_storage] = lambda: storage
    app.dependency_overrides[get_speech_to_text_provider] = lambda: transcription
    app.dependency_overrides[get_video_audio_extractor] = lambda: extractor
    app.dependency_overrides[get_sentence_segmenter] = lambda: segmenter
    try:
        yield authentication_context, runs, tmp_path, transcription, extractor
    finally:
        app.dependency_overrides.pop(get_agent_project_repository, None)
        app.dependency_overrides.pop(get_agent_run_repository, None)
        app.dependency_overrides.pop(get_requirement_source_storage, None)
        app.dependency_overrides.pop(get_speech_to_text_provider, None)
        app.dependency_overrides.pop(get_video_audio_extractor, None)
        app.dependency_overrides.pop(get_sentence_segmenter, None)


def create_project(client):
    response = client.post(
        "/api/v1/agent-projects",
        json={"agent_type": "REQUIREMENTS", "name": "Discovery", "domain": "Health"},
    )
    assert response.status_code == 201
    return response.json()["data"]["agent_project_id"]


def upload(client, project_id, source_type, filename, content):
    response = client.post(
        f"/api/v1/agents/requirements/{project_id}/sources",
        data={"source_type": source_type},
        files={"file": (filename, content, "application/octet-stream")},
    )
    assert response.status_code == 201
    return response.json()["data"]


def process(client, project_id, source_id, *, force=False):
    suffix = "?force=true" if force else ""
    response = client.post(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/processing{suffix}"
    )
    assert response.status_code == 200
    return response.json()["data"]


def prepare(client, project_id, source_id, *, force=False):
    suffix = "?force=true" if force else ""
    return client.post(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/preparation{suffix}"
    )


def prep_status(client, project_id, source_id):
    return client.get(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/preparation/status"
    )


def normalized(client, project_id, source_id):
    return client.get(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/preparation/normalized"
    )


def chunks(client, project_id, source_id):
    return client.get(
        f"/api/v1/agents/requirements/{project_id}/sources/{source_id}/preparation/chunks"
    )


def pdf_bytes(page_texts):
    document = fitz.open()
    for text in page_texts:
        page = document.new_page()
        page.insert_text((72, 72), text)
    value = document.tobytes()
    document.close()
    return value


def docx_bytes():
    output = BytesIO()
    document = Document()
    document.add_paragraph("First paragraph requirement.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Goal"
    document.add_paragraph("Final paragraph constraint.")
    document.save(output)
    return output.getvalue()


def fallback_segmenter(monkeypatch):
    def missing_spacy(_package):
        raise PackageNotFoundError("spacy")

    monkeypatch.setattr(chunker_module, "version", missing_spacy)
    return SentenceSegmenter()


def normalized_evidence(text):
    return NormalizedEvidenceResult(
        source_id="AUD-999",
        agent_project_id="AGP-999",
        source_type=RequirementSourceType.AUDIO,
        source_sha256="a" * 64,
        parent_processor_version="1",
        parent_result_fingerprint="b" * 64,
        transcription_fingerprint="c" * 64,
        normalizer_version="1",
        normalized_text=text,
        provenance=[
            NormalizedProvenanceSpan(
                normalized_start_offset=0,
                normalized_end_offset=len(text),
                parent_start_offset=0,
                parent_end_offset=len(text),
            )
        ],
        created_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
    )


def test_sentence_segmenter_uses_spacy_when_initialization_succeeds(
    monkeypatch,
):
    calls = []

    class FakePipeline:
        def add_pipe(self, name):
            calls.append(("add_pipe", name))

        def __call__(self, text):
            calls.append(("segment", text))
            return SimpleNamespace(
                sents=[
                    SimpleNamespace(start_char=0, end_char=4),
                    SimpleNamespace(start_char=5, end_char=9),
                ]
            )

    def blank(language):
        calls.append(("blank", language))
        return FakePipeline()

    monkeypatch.setattr(chunker_module, "version", lambda _package: "test-spacy")
    monkeypatch.setitem(sys.modules, "spacy", SimpleNamespace(blank=blank))
    segmenter = SentenceSegmenter()

    assert segmenter.identity == ("spacy_blank_en_sentencizer", "test-spacy")
    assert segmenter.segment("One. Two.") == [SentenceSpan(0, 4), SentenceSpan(5, 9)]
    assert calls == [
        ("blank", "en"),
        ("add_pipe", "sentencizer"),
        ("segment", "One. Two."),
    ]


def test_sentence_segmenter_uses_fallback_when_spacy_is_missing(monkeypatch):
    segmenter = fallback_segmenter(monkeypatch)

    assert segmenter.identity == ("deterministic_sentence_fallback", "1")
    assert segmenter.segment("One. Two?") == [SentenceSpan(0, 4), SentenceSpan(5, 9)]


def test_blocked_spacy_import_falls_back_during_audio_preparation(
    preparation_context, monkeypatch
):
    authentication, _, _, _, _ = preparation_context
    segmenter = app.dependency_overrides[get_sentence_segmenter]()
    segmenter._spacy_version = "3.8.16"
    segmenter._pipeline = None

    def blocked_import():
        raise ImportError(
            "DLL load failed while importing align: "
            "An Application Control policy has blocked this file."
        )

    monkeypatch.setattr(segmenter, "_initialize_spacy_pipeline", blocked_import)
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "AUDIO", "interview.mp3", b"ID3\x04\x00\x00audio")
    assert process(client, project_id, source["source_id"])["status"] == "COMPLETED"

    prepared = prepare(client, project_id, source["source_id"]).json()["data"]
    chunk_result = chunks(client, project_id, source["source_id"]).json()["data"]

    assert prepared["status"] == "COMPLETED"
    assert prepared["chunker"]["sentence_segmenter"] == "checking_sentencizer"
    assert chunk_result["chunks"]


def test_sentence_segmenter_falls_back_on_pipeline_initialization_failure(
    monkeypatch,
):
    monkeypatch.setattr(chunker_module, "version", lambda _package: "test-spacy")
    monkeypatch.setattr(
        SentenceSegmenter,
        "_initialize_spacy_pipeline",
        staticmethod(lambda: (_ for _ in ()).throw(RuntimeError("pipeline unavailable"))),
    )
    segmenter = SentenceSegmenter()

    assert segmenter.identity == ("deterministic_sentence_fallback", "1")
    assert segmenter.segment("One. Two.") == [SentenceSpan(0, 4), SentenceSpan(5, 9)]


def test_failed_spacy_initialization_is_cached(monkeypatch):
    attempts = 0

    def fail_once():
        nonlocal attempts
        attempts += 1
        raise ImportError("native extension unavailable")

    monkeypatch.setattr(chunker_module, "version", lambda _package: "test-spacy")
    monkeypatch.setattr(
        SentenceSegmenter,
        "_initialize_spacy_pipeline",
        staticmethod(fail_once),
    )
    segmenter = SentenceSegmenter()

    assert segmenter.segment("One. Two.")
    assert segmenter.segment("Three. Four.")
    assert segmenter.identity == ("deterministic_sentence_fallback", "1")
    assert attempts == 1


def test_sentence_segmenter_does_not_swallow_pipeline_execution_errors(
    monkeypatch,
):
    class BrokenPipeline:
        def __call__(self, _text):
            raise ValueError("invalid pipeline output")

    monkeypatch.setattr(chunker_module, "version", lambda _package: "test-spacy")
    monkeypatch.setattr(
        SentenceSegmenter,
        "_initialize_spacy_pipeline",
        staticmethod(BrokenPipeline),
    )
    segmenter = SentenceSegmenter()

    with pytest.raises(ValueError, match="invalid pipeline output"):
        segmenter.segment("One sentence.")


def test_fallback_chunking_is_deterministic_for_audio_sized_text(monkeypatch):
    text = " ".join(
        f"Stakeholder requirement {index} must remain deterministic."
        for index in range(20)
    )
    normalized = normalized_evidence(text)
    first_segmenter = fallback_segmenter(monkeypatch)
    first_configuration = build_chunker_configuration(
        target_chars=900,
        max_chars=1_200,
        overlap_sentences=1,
        segmenter=first_segmenter,
    )
    first = DeterministicChunker(first_configuration, first_segmenter).chunk(normalized)
    second_segmenter = fallback_segmenter(monkeypatch)
    second_configuration = build_chunker_configuration(
        target_chars=900,
        max_chars=1_200,
        overlap_sentences=1,
        segmenter=second_segmenter,
    )
    second = DeterministicChunker(second_configuration, second_segmenter).chunk(normalized)

    assert first
    assert all(len(chunk.content) <= first_configuration.max_chars for chunk in first)
    assert [chunk.chunk_id for chunk in first] == [chunk.chunk_id for chunk in second]
    assert first_configuration == second_configuration


def test_normalizer_is_conservative_and_deterministic():
    value = "  The client  needs access.  \r\nLine two.\t\r\n\r\n \r\n\r\nEnd.  "
    expected = "The client  needs access.\nLine two.\n\nEnd."

    assert normalize_text(value) == expected
    assert normalize_text(value) == normalize_text(value)
    assert "client  needs" in normalize_text(value)


@pytest.mark.parametrize(
    "filename,content,expected_kind",
    [
        ("requirements.txt", b"  First requirement.\r\n\r\n\r\nSecond requirement.  ", "text"),
        ("brief.pdf", pdf_bytes(["Page one requirement.", "Page two constraint."]), "pages"),
        ("brief.docx", docx_bytes(), "blocks"),
    ],
    ids=["txt", "pdf", "docx"],
)
def test_document_normalization_and_chunk_provenance(
    preparation_context, filename, content, expected_kind
):
    authentication, runs, root, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", filename, content)
    source_snapshot = dict(source)
    assert process(client, project_id, source["source_id"])["status"] == "COMPLETED"

    assert prep_status(client, project_id, source["source_id"]).json()["data"]["status"] == "READY"
    response = prepare(client, project_id, source["source_id"])
    normalized_result = normalized(client, project_id, source["source_id"]).json()["data"]
    chunk_result = chunks(client, project_id, source["source_id"]).json()["data"]

    assert response.status_code == 200
    assert response.json()["data"]["status"] == "COMPLETED"
    assert normalized_result["source_id"] == source["source_id"]
    assert normalized_result["source_sha256"] == source["sha256"]
    assert normalized_result["parent_processor_version"] == "2"
    assert chunk_result["parent_result_fingerprint"] == normalized_result["parent_result_fingerprint"]
    assert all(item["chunk_id"].startswith("CHK-") for item in chunk_result["chunks"])
    for item in chunk_result["chunks"]:
        assert normalized_result["normalized_text"][
            item["normalized_start_offset"] : item["normalized_end_offset"]
        ] == item["content"]
    provenance = normalized_result["provenance"]
    if expected_kind == "pages":
        assert [span["page_numbers"] for span in provenance] == [[1], [2]]
    elif expected_kind == "blocks":
        assert [span["document_blocks"][0]["block_type"] for span in provenance] == [
            "PARAGRAPH",
            "TABLE",
            "PARAGRAPH",
        ]
        assert [span["document_blocks"][0]["block_index"] for span in provenance] == [0, 1, 2]
    else:
        assert provenance[0]["parent_start_offset"] == 0
        assert provenance[0]["parent_end_offset"] > 0
    metadata = client.get(
        f"/api/v1/agents/requirements/{project_id}/sources"
    ).json()["data"][0]
    assert metadata == source_snapshot
    assert (root / project_id / "inputs" / f"{source['source_id']}.normalized-result.json").is_file()
    assert (root / project_id / "inputs" / f"{source['source_id']}.chunk-result.json").is_file()
    assert runs.documents == {}


@pytest.mark.parametrize(
    "source_type,filename,content,source_prefix",
    [
        ("AUDIO", "interview.mp3", b"ID3\x04\x00\x00interview", "AUD-"),
        ("VIDEO", "workflow.mp4", b"\x00\x00\x00\x18ftypisomvideo", "VID-"),
    ],
)
def test_media_chunks_preserve_segments_timestamps_and_fingerprints(
    preparation_context, source_type, filename, content, source_prefix
):
    authentication, _, _, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, source_type, filename, content)
    assert source["source_id"].startswith(source_prefix)
    assert process(client, project_id, source["source_id"])["status"] == "COMPLETED"
    assert prepare(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"

    normalized_result = normalized(client, project_id, source["source_id"]).json()["data"]
    chunk_result = chunks(client, project_id, source["source_id"]).json()["data"]
    references = [
        reference
        for span in normalized_result["provenance"]
        for reference in span["media_segments"]
    ]
    assert [reference["segment_index"] for reference in references] == [0, 1]
    assert [(reference["start_seconds"], reference["end_seconds"]) for reference in references] == [
        (0.0, 1.5),
        (1.5, 3.0),
    ]
    assert normalized_result["transcription_fingerprint"] == "a" * 64
    if source_type == "VIDEO":
        assert normalized_result["extraction_fingerprint"] == "b" * 64
    else:
        assert normalized_result["extraction_fingerprint"] is None
    assert all(chunk["provenance"][0]["media_segments"] for chunk in chunk_result["chunks"])


def test_reuse_force_and_deterministic_ids(preparation_context):
    authentication, _, _, transcription, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "AUDIO", "interview.mp3", b"ID3\x04\x00\x00audio")
    process(client, project_id, source["source_id"])

    prepare(client, project_id, source["source_id"])
    first = chunks(client, project_id, source["source_id"]).json()["data"]
    prepare(client, project_id, source["source_id"])
    reused = chunks(client, project_id, source["source_id"]).json()["data"]
    prepare(client, project_id, source["source_id"], force=True)
    forced = chunks(client, project_id, source["source_id"]).json()["data"]

    assert [chunk["chunk_id"] for chunk in first["chunks"]] == [
        chunk["chunk_id"] for chunk in reused["chunks"]
    ] == [chunk["chunk_id"] for chunk in forced["chunks"]]
    assert transcription.calls == 1


def test_changed_parent_model_fingerprint_invalidates_chunks(preparation_context):
    authentication, _, _, transcription, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "AUDIO", "interview.mp3", b"ID3\x04\x00\x00audio")
    process(client, project_id, source["source_id"])
    prepare(client, project_id, source["source_id"])
    first = chunks(client, project_id, source["source_id"]).json()["data"]

    transcription.configuration = SpeechToTextConfiguration(
        provider="fake",
        provider_version="test-2",
        model_name="fake-medium",
        device="cpu",
        compute_type="int8",
        language=None,
        fingerprint="c" * 64,
    )
    assert process(client, project_id, source["source_id"], force=True)["status"] == "COMPLETED"
    assert prep_status(client, project_id, source["source_id"]).json()["data"]["status"] == "READY"
    prepare(client, project_id, source["source_id"])
    second = chunks(client, project_id, source["source_id"]).json()["data"]

    assert second["parent_result_fingerprint"] != first["parent_result_fingerprint"]
    assert [chunk["chunk_id"] for chunk in second["chunks"]] != [
        chunk["chunk_id"] for chunk in first["chunks"]
    ]


def test_incompatible_chunk_config_changes_identity(preparation_context):
    authentication, _, _, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", "requirements.txt", b"One. Two. Three.")
    process(client, project_id, source["source_id"])
    prepare(client, project_id, source["source_id"])
    value = NormalizedEvidenceResult.model_validate(
        normalized(client, project_id, source["source_id"]).json()["data"]
    )
    segmenter = SentenceSegmenter()
    first = DeterministicChunker(
        build_chunker_configuration(
            target_chars=100, max_chars=120, overlap_sentences=0, segmenter=segmenter
        ),
        segmenter,
    ).chunk(value)
    second = DeterministicChunker(
        build_chunker_configuration(
            target_chars=100, max_chars=120, overlap_sentences=1, segmenter=segmenter
        ),
        segmenter,
    ).chunk(value)

    assert [chunk.chunk_id for chunk in first] != [chunk.chunk_id for chunk in second]


def test_parent_must_be_completed_and_valid(preparation_context):
    authentication, _, root, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", "requirements.txt", b"A valid requirement.")

    rejected = prepare(client, project_id, source["source_id"]).json()["data"]
    assert rejected["status"] == "FAILED"
    assert rejected["error"]["code"] == "PARENT_PROCESSING_NOT_COMPLETED"
    assert normalized(client, project_id, source["source_id"]).status_code == 404

    process(client, project_id, source["source_id"])
    result_path = root / project_id / "inputs" / "DOC-001.document-result.json"
    result_path.write_text("not-json", encoding="utf-8")
    rejected = prepare(client, project_id, source["source_id"]).json()["data"]
    assert rejected["status"] == "FAILED"
    assert rejected["error"]["code"] == "PARENT_PROCESSING_NOT_COMPLETED"
    assert "not-json" not in json.dumps(rejected)


def test_empty_normalized_text_is_rejected_safely(preparation_context):
    authentication, _, root, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", "requirements.txt", b"Original")
    process(client, project_id, source["source_id"])
    result_path = root / project_id / "inputs" / "DOC-001.document-result.json"
    parent = json.loads(result_path.read_text(encoding="utf-8"))
    parent["text"] = "   "
    parent["structural_provenance"]["text_ranges"] = [
        {"text": "   ", "start_offset": 0, "end_offset": 3}
    ]
    result_path.write_text(json.dumps(parent), encoding="utf-8")

    response = prepare(client, project_id, source["source_id"]).json()["data"]

    assert response["status"] == "FAILED"
    assert response["error"]["code"] == "NORMALIZED_TEXT_EMPTY"
    assert normalized(client, project_id, source["source_id"]).status_code == 404


def test_normalizer_version_invalidates_cached_artifacts(
    preparation_context, monkeypatch
):
    from app.modules.agents.requirements.processing.preparation import service as preparation_module

    authentication, _, _, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", "requirements.txt", b"Versioned requirement.")
    process(client, project_id, source["source_id"])
    prepare(client, project_id, source["source_id"])
    first_ids = [
        chunk["chunk_id"]
        for chunk in chunks(client, project_id, source["source_id"]).json()["data"]["chunks"]
    ]

    monkeypatch.setattr(preparation_module, "NORMALIZER_VERSION", "99")
    assert prep_status(client, project_id, source["source_id"]).json()["data"]["status"] == "READY"
    assert prepare(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"
    second_ids = [
        chunk["chunk_id"]
        for chunk in chunks(client, project_id, source["source_id"]).json()["data"]["chunks"]
    ]

    assert second_ids != first_ids


def test_preparation_enforces_owner_type_and_archive_guards(preparation_context):
    authentication, _, _, _, _ = preparation_context
    owner = authentication.client("USR-001")
    foreign = authentication.client("USR-002")
    project_id = create_project(owner)
    source = upload(owner, project_id, "DOCUMENT", "requirements.txt", b"Requirement")
    process(owner, project_id, source["source_id"])

    assert prepare(foreign, project_id, source["source_id"]).status_code == 404
    wrong_project = owner.post(
        "/api/v1/agent-projects",
        json={"agent_type": "IMPLEMENTATION", "name": "Wrong", "domain": "Health"},
    ).json()["data"]["agent_project_id"]
    assert prep_status(owner, wrong_project, "DOC-001").status_code == 422
    owner.patch(f"/api/v1/agent-projects/{project_id}", json={"status": "ARCHIVED"})
    assert prepare(owner, project_id, source["source_id"]).status_code == 409
    assert prep_status(owner, project_id, source["source_id"]).status_code == 409


def test_missing_or_corrupt_preparation_artifact_never_remains_completed(
    preparation_context,
):
    authentication, _, root, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", "requirements.txt", b"One requirement.")
    process(client, project_id, source["source_id"])
    prepare(client, project_id, source["source_id"])
    directory = root / project_id / "inputs"

    (directory / "DOC-001.chunk-result.json").unlink()
    reconciled = prep_status(client, project_id, source["source_id"]).json()["data"]
    assert reconciled["status"] == "FAILED"
    assert reconciled["error"]["code"] == "PREPARATION_RESULT_INVALID"
    assert chunks(client, project_id, source["source_id"]).status_code == 404

    assert prepare(client, project_id, source["source_id"]).json()["data"]["status"] == "COMPLETED"
    (directory / "DOC-001.normalized-result.json").write_text("{bad", encoding="utf-8")
    assert normalized(client, project_id, source["source_id"]).status_code == 404
    assert prep_status(client, project_id, source["source_id"]).json()["data"]["status"] == "FAILED"


def test_source_deletion_removes_preparation_artifacts(preparation_context):
    authentication, _, root, _, _ = preparation_context
    client = authentication.client()
    project_id = create_project(client)
    source = upload(client, project_id, "DOCUMENT", "requirements.txt", b"One requirement.")
    process(client, project_id, source["source_id"])
    prepare(client, project_id, source["source_id"])
    directory = root / project_id / "inputs"

    response = client.delete(
        f"/api/v1/agents/requirements/{project_id}/sources/{source['source_id']}"
    )

    assert response.status_code == 200
    assert not list(directory.glob("DOC-001.*"))
