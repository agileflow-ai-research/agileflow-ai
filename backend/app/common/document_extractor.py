import hashlib
import io
import mimetypes
from pathlib import Path
from typing import Any

import docx
import pypdf

MAX_DOCUMENT_SIZE_BYTES = 10 * 1024 * 1024  # 10 MiB

ALLOWED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md"}

ALLOWED_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/pdf",
    "text/plain",
    "text/markdown",
    "text/x-markdown",
    "application/octet-stream",
    "application/zip",
}


class DocumentExtractorError(ValueError):
    """Base error for document extraction failures."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class UnsupportedDocumentTypeError(DocumentExtractorError):
    def __init__(
        self,
        message: str = "Unsupported document format. Supported formats are: .docx, .pdf, .txt, .md",
    ) -> None:
        super().__init__("UNSUPPORTED_DOCUMENT_TYPE", message)


class DocumentTooLargeError(DocumentExtractorError):
    def __init__(
        self,
        message: str = "Document exceeds the maximum allowed size of 10 MiB",
    ) -> None:
        super().__init__("DOCUMENT_TOO_LARGE", message)


class DocumentTextNotFoundError(DocumentExtractorError):
    def __init__(
        self,
        message: str = "No extractable text found in the uploaded document. Scanned or image-only documents are not supported.",
    ) -> None:
        super().__init__("DOCUMENT_TEXT_NOT_FOUND", message)


class DocumentExtractionFailedError(DocumentExtractorError):
    def __init__(
        self,
        message: str = "Document extraction failed unexpectedly",
    ) -> None:
        super().__init__("DOCUMENT_EXTRACTION_FAILED", message)


def _extract_docx(content: bytes) -> str:
    try:
        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        table_rows: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    deduped: list[str] = []
                    for cell in cells:
                        if not deduped or deduped[-1] != cell:
                            deduped.append(cell)
                    table_rows.append(" | ".join(deduped))
        parts = paragraphs
        if table_rows:
            parts = parts + ["\n--- Table Content ---\n" + "\n".join(table_rows)]
        return "\n\n".join(parts)
    except Exception as exc:
        raise DocumentExtractionFailedError(f"Failed to extract DOCX document: {exc}") from exc


def _extract_pdf(content: bytes) -> str:
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        pages_text: list[str] = []
        for index, page in enumerate(reader.pages):
            extracted = page.extract_text() or ""
            cleaned = extracted.strip()
            if cleaned:
                pages_text.append(cleaned)
        return "\n\n".join(pages_text)
    except Exception as exc:
        raise DocumentExtractionFailedError(f"Failed to extract PDF document: {exc}") from exc


def _extract_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionFailedError("Failed to decode text file using supported encodings")


def extract_document_text(
    content: bytes,
    filename: str,
    content_type: str | None = None,
) -> tuple[str, str, int, str]:
    """
    Extracts text from an uploaded document.
    Validates extension, size, MIME type, extracts content, and trims outer whitespace.

    Returns:
        tuple of (extracted_text, detected_mime_type, character_count, sha256_hex)
    """
    if len(content) > MAX_DOCUMENT_SIZE_BYTES:
        raise DocumentTooLargeError()

    if not content or len(content) == 0:
        raise DocumentTextNotFoundError("Uploaded document is empty")

    clean_filename = Path(filename).name
    ext = Path(clean_filename).suffix.lower()

    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedDocumentTypeError(
            f"Unsupported file extension '{ext}'. Supported formats are: .docx, .pdf, .txt, .md"
        )

    # Determine / validate MIME type
    guessed_type, _ = mimetypes.guess_type(clean_filename)
    detected_mime = content_type or guessed_type or "application/octet-stream"

    if content_type and content_type.lower() not in ALLOWED_MIME_TYPES:
        # If client sent an unsupported MIME type and it doesn't match extension
        if not (ext in {".txt", ".md"} and content_type.startswith("text/")):
            raise UnsupportedDocumentTypeError(
                f"Unsupported content type '{content_type}' for document '{clean_filename}'"
            )

    sha256_hash = hashlib.sha256(content).hexdigest()

    if ext == ".docx":
        raw_text = _extract_docx(content)
        final_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == ".pdf":
        raw_text = _extract_pdf(content)
        final_mime = "application/pdf"
    elif ext in {".txt", ".md"}:
        raw_text = _extract_text(content)
        final_mime = "text/markdown" if ext == ".md" else "text/plain"
    else:
        raise UnsupportedDocumentTypeError(f"Unsupported extension '{ext}'")

    normalized_text = raw_text.strip()
    if not normalized_text:
        raise DocumentTextNotFoundError()

    return normalized_text, final_mime, len(normalized_text), sha256_hash
