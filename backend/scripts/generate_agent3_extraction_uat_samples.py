"""Generate small, non-sensitive documents for Agent 3 extraction UAT."""

import argparse
import json
from pathlib import Path

import fitz
import yaml
from docx import Document


def create_samples(output_directory: Path) -> list[Path]:
    output_directory.mkdir(parents=True, exist_ok=True)
    created: list[Path] = []

    pdf_path = output_directory / "api-specification.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "API Specification\nPOST /bookings creates a booking.")
    second_page = pdf.new_page()
    second_page.insert_text((72, 72), "GET /bookings/{id} returns a booking.")
    pdf.save(pdf_path)
    pdf.close()
    created.append(pdf_path)

    docx_path = output_directory / "implementation-plan.docx"
    document = Document()
    document.add_heading("Implementation Plan", level=1)
    document.add_paragraph("Build the booking API and persistence service.")
    table = document.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Method"
    table.rows[0].cells[1].text = "Endpoint"
    table.rows[0].cells[2].text = "Purpose"
    cells = table.add_row().cells
    cells[0].text = "POST"
    cells[1].text = "/bookings"
    cells[2].text = "Create booking"
    document.save(docx_path)
    created.append(docx_path)

    text_samples = {
        "requirements.txt": "Users can create and cancel bookings.\n",
        "requirements.md": "# Requirements\n\n- Create a booking\n- Cancel a booking\n",
        "api.json": json.dumps(
            {"endpoint": "/bookings", "method": "POST"}, indent=2, sort_keys=True
        )
        + "\n",
        "architecture.yaml": yaml.safe_dump(
            {"frontend": "React", "backend": "Express"}, sort_keys=True
        ),
    }
    for filename, content in text_samples.items():
        path = output_directory / filename
        path.write_text(content, encoding="utf-8")
        created.append(path)

    scanned_path = output_directory / "scanned-no-text.pdf"
    scanned = fitz.open()
    page = scanned.new_page()
    pixels = fitz.Pixmap(fitz.csRGB, (0, 0, 200, 100), False)
    pixels.clear_with(235)
    page.insert_image(page.rect, stream=pixels.tobytes("png"))
    scanned.save(scanned_path)
    scanned.close()
    created.append(scanned_path)
    return created


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output-dir",
        type=Path,
        required=True,
        help="Directory in which to create the disposable UAT documents.",
    )
    arguments = parser.parse_args()
    for path in create_samples(arguments.output_dir.resolve()):
        print(path)


if __name__ == "__main__":
    main()
